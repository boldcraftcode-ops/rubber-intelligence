"""API unggah file data lewat web (Fase 6).

- Hanya .xlsx, dengan batas ukuran (UNGGAH_MAKS_MB, default 20 MB).
- Format mengikuti file yang sudah dipakai pengguna:
  * Ekspor "Master Data" aplikasi internal (tiket intake), ATAU
  * Format template (sheet Pembelian_Harian / Stok_Harian / SICOM_Harian /
    Kurs / Parameter_Pabrik / Target_Margin / Perencanaan / Master_Vendor /
    Master_Grade / Intake_Tiket) — boleh hanya berisi sebagian sheet.
- Jenis file dideteksi otomatis; file asli disimpan apa adanya di folder
  `unggahan/` (ikut backup harian) dan isi barisnya di lapisan raw.
- Idempoten: file yang sama diunggah dua kali tidak menggandakan data.
"""
from __future__ import annotations

import datetime as dt
import os
import pathlib
import re

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session

from .core.alarm import catat_audit
from .importer.template import impor_master_data, impor_template
from .importer.tiket import cari_header

SHEET_TEMPLATE = {
    "Pembelian_Harian", "SICOM_Harian", "Kurs", "Stok_Harian",
    "Parameter_Pabrik", "Aturan_Stok_Harga", "Target_Margin", "Perencanaan",
    "Master_Vendor", "Master_Grade", "Intake_Tiket", "Harga_Grade",
}

BATAS_PERINGATAN = 200  # maksimal peringatan yang dikirim ke layar


def _nama_aman(nama: str) -> str:
    dasar = pathlib.Path(nama).name
    return re.sub(r"[^A-Za-z0-9._ -]", "_", dasar) or "unggahan.xlsx"


def deteksi_jenis(path: str) -> str | None:
    """'template' | 'tiket' | None (tak dikenal)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if set(wb.sheetnames) & SHEET_TEMPLATE:
            return "template"
        for nama in wb.sheetnames:
            if cari_header(wb[nama]) is not None:
                return "tiket"
        return None
    finally:
        wb.close()


def pasang_router_unggah(app, ambil_session):
    router = APIRouter(prefix="/api")
    dep = Depends(ambil_session)

    @router.post("/unggah")
    async def unggah(file: UploadFile = File(...), pengisi: str = Form("?"),
                     session: Session = dep):
        nama = _nama_aman(file.filename or "unggahan.xlsx")
        if not nama.lower().endswith(".xlsx"):
            raise HTTPException(422, f"File '{nama}' bukan .xlsx. Hanya file Excel "
                                     ".xlsx yang diterima — simpan/ekspor ulang "
                                     "sebagai .xlsx lalu unggah lagi.")
        isi = await file.read()
        maks_mb = float(os.environ.get("UNGGAH_MAKS_MB", "20"))
        if len(isi) > maks_mb * 1024 * 1024:
            raise HTTPException(413, f"File {len(isi) / 1024 / 1024:.1f} MB melebihi "
                                     f"batas {maks_mb:.0f} MB. Pecah filenya "
                                     "(mis. per tahun) lalu unggah bertahap.")
        if not isi:
            raise HTTPException(422, "File kosong. Pilih file yang benar.")

        # Simpan file asli apa adanya (ikut backup harian; tidak pernah dihapus)
        folder = pathlib.Path(os.environ.get("UNGGAH_DIR", "unggahan"))
        folder.mkdir(parents=True, exist_ok=True)
        stempel = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        path = folder / f"{stempel}_{nama}"
        path.write_bytes(isi)

        try:
            jenis = deteksi_jenis(str(path))
        except Exception:
            raise HTTPException(422, f"File '{nama}' tidak bisa dibuka sebagai "
                                     "Excel. Pastikan file tidak rusak/terkunci "
                                     "password, lalu coba lagi.")
        if jenis is None:
            raise HTTPException(422, f"Format '{nama}' tidak dikenal. Yang diterima: "
                                     "(1) ekspor 'Master Data' tiket dari aplikasi "
                                     "internal, atau (2) file dengan sheet format "
                                     "template: " + ", ".join(sorted(SHEET_TEMPLATE)) + ".")

        fn = impor_master_data if jenis == "tiket" else impor_template
        hasil = fn(str(path), session, pengunggah=pengisi)

        catat_audit(session, pengisi, "unggah", "file", {
            "nama_file": nama, "jenis": jenis, "ukuran_kb": round(len(isi) / 1024, 1),
            "unggahan_id": hasil.unggahan_id, "jumlah": hasil.jumlah,
            "ditolak": hasil.ditolak, "duplikat_file": hasil.duplikat_file,
        })
        session.commit()

        return {
            "status": "selesai",
            "jenis": "Tiket intake (Master Data)" if jenis == "tiket"
                     else "Data berkala (format template)",
            "nama_file": nama,
            "unggahan_id": hasil.unggahan_id,
            "duplikat_file": hasil.duplikat_file,
            "jumlah": hasil.jumlah,
            "ditolak": hasil.ditolak,
            "jumlah_peringatan": len(hasil.peringatan),
            "peringatan": hasil.peringatan[:BATAS_PERINGATAN],
            "pesan": ("File sama persis pernah diunggah — data tidak digandakan "
                      "(aman)." if hasil.duplikat_file else
                      "Impor selesai; mart & alarm sudah dihitung ulang otomatis."),
        }

    app.include_router(router)
