"""Tes halaman "Stok vs Margin" versi yang mudah dibaca (v0.9.5).

Latar belakang: satu angka korelasi untuk seluruh periode menyesatkan karena
tahun yang berbeda bisa punya hubungan berbeda arah. Tes di sini menjaga agar
API selalu mengirim bahan untuk menjelaskannya:

- `per_tahun`  : korelasi & rata-rata tiap tahun (2023 timpang jauh dari 2024);
- `kelompok`   : margin rata-rata per kelompok rasio stok (tipis/sedang/tebal);
- `peringatan` : kalimat siap tampil bila hubungan tidak stabil;
- `tahun`      : menyaring titik yang ditampilkan tanpa mengubah per_tahun;
- ekspor `.xlsx` dengan dua sheet (Data bulanan + Ringkasan).
"""
import openpyxl
import pytest

from fastapi.testclient import TestClient

from app.main import buat_app

URL = "/api/analisis/stok-margin"


@pytest.fixture(scope="module")
def klien(engine):
    return TestClient(buat_app(engine))


def test_ringkasan_dasar_tetap_ada(klien):
    """Kunci lama tidak boleh hilang: dipakai layar & tes lain."""
    d = klien.get(URL).json()
    assert d["jumlah_bulan"] >= 40
    assert 0.2 < d["korelasi"] < 0.7
    assert d["margin_stok_rendah"] < d["margin_stok_tinggi"]
    assert "bukan bukti sebab-akibat" in d["catatan"]
    assert d["selisih_tinggi_rendah"] == pytest.approx(
        d["margin_stok_tinggi"] - d["margin_stok_rendah"], abs=1e-4)


def test_korelasi_per_tahun_dan_peringatan_ketidakstabilan(klien):
    d = klien.get(URL).json()
    per_tahun = {p["tahun"]: p for p in d["per_tahun"]}
    assert sorted(per_tahun) == [2023, 2024, 2025, 2026]
    assert per_tahun[2024]["korelasi"] > 0.7      # hubungan kuat
    assert per_tahun[2023]["korelasi"] < 0.3      # nyaris tidak ada
    # Bukti bahwa satu angka gabungan memang mencampur rezim berbeda:
    assert d["peringatan"], "ketidakstabilan antar tahun wajib diberi tahu"
    assert "TIDAK stabil" in d["peringatan"][0]
    assert "2024" in d["peringatan"][0]


def test_kelompok_stok_menutup_semua_bulan(klien):
    d = klien.get(URL).json()
    assert len(d["kelompok"]) == 3
    assert sum(k["jumlah_bulan"] for k in d["kelompok"]) == d["jumlah_bulan"]
    tipis, sedang, tebal = d["kelompok"]
    assert tipis["margin_rata"] is not None and tebal["margin_rata"] is not None
    # Definisi ambang dipegang: <0,5 / 0,5–0,8 / >=0,8
    assert (tipis["batas_atas"], sedang["batas_bawah"], sedang["batas_atas"],
            tebal["batas_bawah"]) == (0.5, 0.5, 0.8, 0.8)


def test_filter_tahun_hanya_menyaring_tampilan(klien):
    semua = klien.get(URL).json()
    satu = klien.get(URL, params={"tahun": 2025}).json()
    assert satu["tahun"] == 2025
    assert {t["tahun"] for t in satu["titik"]} == {2025}
    assert satu["jumlah_bulan"] == len(satu["titik"]) < semua["jumlah_bulan"]
    assert satu["korelasi"] == pytest.approx(
        next(p["korelasi"] for p in semua["per_tahun"] if p["tahun"] == 2025),
        abs=1e-4)
    # perbandingan antar tahun tetap lengkap walau tampilan disaring
    assert len(satu["per_tahun"]) == len(semua["per_tahun"])


def test_tahun_tanpa_data_tidak_galat(klien):
    d = klien.get(URL, params={"tahun": 1999}).json()
    assert d["jumlah_bulan"] == 0
    assert d["titik"] == []
    assert d["korelasi"] is None
    assert any("kurang dari 3" in p for p in d["peringatan"])


def test_ekspor_xlsx_dua_sheet(klien):
    r = klien.get(URL, params={"format": "xlsx", "tahun": 2026})
    assert r.status_code == 200
    assert "stok_vs_margin_2026.xlsx" in r.headers["content-disposition"]
    import io
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Data bulanan", "Ringkasan"]
    ws = wb["Data bulanan"]
    assert [c.value for c in ws[1]] == ["Tahun", "Bulan", "Rasio stok / safety",
                                        "OP Margin (USc/kg)", "Ton"]
    assert ws.max_row == 10          # header + 9 bulan 2026
    ringkas = {b[0].value: b[1].value for b in wb["Ringkasan"]["A2:B30"] if b[0].value}
    assert ringkas["Bulan berdata"] == 9
