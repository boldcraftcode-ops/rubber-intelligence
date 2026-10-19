"""Tes Fase 2: kriteria penerimaan importer, pencarian kurs, penolakan baris,
Qty dari tiket, dan API input cepat SICOM/kurs."""
import datetime as dt

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.lookup import kurs_pada
from app.importer.template import impor_template
from app.main import buat_app
from app.models import Base, Kurs, Pembelian, Sicom

from tests.conftest import TEMPLATE


# ----------------------------------------------------------------------------
# Kriteria 1 & 5: jumlah baris data awal yang terimpor
# ----------------------------------------------------------------------------
JUMLAH_DIHARAPKAN = {
    "pembelian": 2179,
    "sicom": 2441,
    "kurs": 85,
    "stok": 1347,
    "parameter_pabrik": 6,
    "target_margin": 82,
    "perencanaan": 12,
    "vendor": 48,
    "grade": 9,
}


def test_jumlah_baris_data_awal(tmp_path):
    eng = create_engine(f"sqlite:///{tmp_path/'awal.db'}", future=True)
    Base.metadata.create_all(eng)
    Session = sessionmaker(bind=eng, future=True)
    with Session() as s:
        hasil = impor_template(str(TEMPLATE), s, pengunggah="pytest")
    for kunci, diharapkan in JUMLAH_DIHARAPKAN.items():
        assert hasil.jumlah[kunci] == diharapkan, \
            f"{kunci}: terimpor {hasil.jumlah[kunci]}, diharapkan {diharapkan}"


# ----------------------------------------------------------------------------
# Kriteria 2: kurs yang berlaku (baris terbaru dengan berlaku_dari <= tanggal)
# ----------------------------------------------------------------------------
@pytest.mark.parametrize("tanggal,nilai", [
    (dt.date(2020, 3, 21), 15055),
    (dt.date(2020, 3, 25), 16342),
    (dt.date(2020, 3, 31), 16263),
    (dt.date(2020, 4, 1), 16367),
])
def test_kurs_berlaku(session, tanggal, nilai):
    k = kurs_pada(session, tanggal)
    assert k is not None
    assert float(k.idr_per_usd) == nilai


# ----------------------------------------------------------------------------
# Kriteria 3: baris salah ditolak dengan pesan jelas; unggah ulang aman
# ----------------------------------------------------------------------------
def _buat_file_rusak(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pembelian_Harian"
    ws.append(["Tanggal", "Item_Code", "Kode_Vendor", "Ticket_No",
               "Qty_Kering_Kg", "Amount_Rp"])
    ws.append([dt.date(2026, 10, 16), "ALL", None, None, 1000, 25_000_000])  # baik
    ws.append(["bukan-tanggal", "ALL", None, None, 1000, 25_000_000])        # tanggal salah
    ws.append([dt.date(2026, 10, 17), "ALL", None, None, "abc", 25_000_000])  # qty salah
    ws.append([dt.date(2026, 10, 18), "ALL", None, None, 1000, -5])           # amount negatif
    ws2 = wb.create_sheet("SICOM_Harian")
    ws2.append(["Tanggal", "SICOM_USc_per_kg"])
    ws2.append([dt.date(2026, 10, 16), 250.5])   # baik
    ws2.append([dt.date(2026, 10, 17), "x"])     # angka salah
    ws2.append([None, 251.0])                    # tanggal kosong
    ws3 = wb.create_sheet("Kurs")
    ws3.append(["Berlaku_Dari", "Kurs_IDR_per_USD"])
    ws3.append([dt.date(2026, 10, 16), 0])       # nol -> ditolak
    wb.save(path)


def test_baris_salah_ditolak(tmp_path):
    eng = create_engine(f"sqlite:///{tmp_path/'rusak.db'}", future=True)
    Base.metadata.create_all(eng)
    Session = sessionmaker(bind=eng, future=True)
    f = tmp_path / "rusak.xlsx"
    _buat_file_rusak(f)
    with Session() as s:
        hasil = impor_template(str(f), s, pengunggah="pytest")
        assert hasil.jumlah["pembelian"] == 1   # hanya baris baik
        assert hasil.jumlah["sicom"] == 1
        assert hasil.jumlah["kurs"] == 0
        assert hasil.ditolak["Pembelian_Harian"] == 3
        assert hasil.ditolak["SICOM_Harian"] == 2
        assert hasil.ditolak["Kurs"] == 1
        teks = "\n".join(hasil.peringatan)
        assert "bukan-tanggal" in teks and "dd/mm/yyyy" in teks
        assert "'abc'" in teks and "tidak terbaca" in teks
        assert "DITOLAK" in teks
        # unggah ulang file yang sama: tidak menggandakan
        hasil2 = impor_template(str(f), s, pengunggah="pytest")
        assert hasil2.duplikat_file
        assert s.scalar(select(func.count()).select_from(Pembelian)) == 1
        assert s.scalar(select(func.count()).select_from(Sicom)) == 1


# ----------------------------------------------------------------------------
# Kriteria: Qty_Kering_Kg kosong dihitung dari tiket (tingkat terhalus)
# ----------------------------------------------------------------------------
def _buat_file_qty_dari_tiket(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Intake_Tiket"
    ws.append(["Ticket Date", "Ticket No", "Supplier Code", "Supplier Name",
               "Vehicle No", "Item Code", "DRC Date", "Wet", "Sample", "Reject",
               "DRC", "Dry", "Asal Getah", "Provinsi", "Kabupaten/Kota",
               "Kecamatan", "Desa", "BIN", "Location"])
    t = dt.date(2026, 10, 16)
    ws.append([t, "T001", "V1", "VENDOR SATU", "BK 1", "SLAB2", t,
               10000, 10, 0, 60, 6000, "SUMUT, A, B, C", "SUMUT", "A", "B", "C",
               "BIN1", "L1"])
    ws.append([t, "T002", "V2", "VENDOR DUA", "BK 2", "SLAB1", t,
               2000, 5, 0, 55, 1100, "SUMUT, A, B, C", "SUMUT", "A", "B", "C",
               "BIN1", "L1"])
    ws2 = wb.create_sheet("Pembelian_Harian")
    ws2.append(["Tanggal", "Item_Code", "Kode_Vendor", "Ticket_No",
                "Qty_Kering_Kg", "Amount_Rp"])
    # Qty kosong: per tiket, per grade, dan total — semua dihitung dari tiket
    ws2.append([t, "SLAB2", None, "T001", None, 180_000_000])
    ws2.append([dt.date(2026, 10, 17), "ALL", None, None, None, 1_000])  # tak ada tiket -> tolak
    wb.save(path)


def test_qty_dari_tiket(tmp_path):
    eng = create_engine(f"sqlite:///{tmp_path/'qty.db'}", future=True)
    Base.metadata.create_all(eng)
    Session = sessionmaker(bind=eng, future=True)
    f = tmp_path / "qty.xlsx"
    _buat_file_qty_dari_tiket(f)
    with Session() as s:
        hasil = impor_template(str(f), s, pengunggah="pytest")
        assert hasil.jumlah["tiket_baris"] == 2
        assert hasil.jumlah["pembelian"] == 1
        assert hasil.ditolak["Pembelian_Harian"] == 1  # 17 Okt tanpa tiket
        p = s.scalar(select(Pembelian).where(Pembelian.ticket_no == "T001"))
        assert p is not None and float(p.qty_kering_kg) == 6000  # Dry tiket T001
        teks = "\n".join(hasil.peringatan)
        assert "dipakai jumlah Dry tiket" in teks
        assert "Isi Qty_Kering_Kg atau unggah tiketnya dahulu" in teks


# ----------------------------------------------------------------------------
# Kriteria 4: API input cepat SICOM & kurs (form + riwayat)
# ----------------------------------------------------------------------------
@pytest.fixture()
def klien():
    eng = create_engine("sqlite://", future=True,
                        connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    return TestClient(buat_app(eng))


def test_api_sicom_simpan_dan_riwayat(klien):
    r = klien.post("/api/sicom", json={"tanggal": "2026-10-16",
                                       "usc_per_kg": 259.3, "pengisi": "tes"})
    assert r.status_code == 200
    assert r.json()["pesan"] == "Harga SICOM baru tersimpan."
    # timpa nilai yang sama tanggalnya -> tidak berganda, ada pesan nilai lama
    r = klien.post("/api/sicom", json={"tanggal": "2026-10-16",
                                       "usc_per_kg": 260.0, "pengisi": "tes"})
    assert "Nilai lama 259.30 diganti menjadi 260.00." == r.json()["pesan"]
    riwayat = klien.get("/api/sicom").json()
    assert len(riwayat) == 1
    assert riwayat[0]["usc_per_kg"] == 260.0


def test_api_sicom_validasi(klien):
    r = klien.post("/api/sicom", json={"tanggal": "2026-10-16", "usc_per_kg": -5})
    assert r.status_code == 422
    assert "lebih dari 0" in r.json()["detail"]
    r = klien.post("/api/sicom", json={"tanggal": "bukan", "usc_per_kg": 100})
    assert r.status_code == 422
    assert "tanggal tidak valid" in r.json()["detail"].lower()


def test_api_kurs_dan_berlaku(klien):
    klien.post("/api/kurs", json={"berlaku_dari": "2020-03-20", "idr_per_usd": 15055})
    klien.post("/api/kurs", json={"berlaku_dari": "2020-03-24", "idr_per_usd": 16342})
    r = klien.get("/api/kurs/berlaku", params={"tanggal": "2020-03-21"})
    assert r.json()["idr_per_usd"] == 15055
    r = klien.get("/api/kurs/berlaku", params={"tanggal": "2020-03-25"})
    assert r.json()["idr_per_usd"] == 16342
    r = klien.get("/api/kurs/berlaku", params={"tanggal": "2020-01-01"})
    assert r.status_code == 404
    assert "Belum ada kurs" in r.json()["detail"]
    assert len(klien.get("/api/kurs").json()) == 2
