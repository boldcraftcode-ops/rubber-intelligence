"""Tes #27: pratinjau sebelum masuk database + unduh Excel data mismatch.

Alur yang diuji meniru layar "Unggah Data":
pilih file -> pratinjau (tidak ada yang tersimpan) -> unduh mismatch ->
konfirmasi (baru tersimpan) / batalkan.
"""
from __future__ import annotations

import datetime as dt
import io
import pathlib

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import buat_app
from app.models import PeringatanData, RawBaris, RawUnggahan, Stok


@pytest.fixture()
def klien(tmp_path, monkeypatch):
    monkeypatch.setenv("UNGGAH_DIR", str(tmp_path / "unggahan"))
    eng = create_engine("sqlite://", future=True,
                        connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    Sesi = sessionmaker(bind=eng, future=True)
    yield TestClient(buat_app(eng)), Sesi, tmp_path / "unggahan"


def _xlsx_stok(dengan_mismatch: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stok_Harian"
    ws.append(["Tanggal", "Stok_Kg", "Item_Code", "Catatan"])
    ws.append([dt.date(2026, 10, 20), 500000, None, "contoh baik"])
    if dengan_mismatch:
        ws.append([None, 100, None, "tanggal lupa diisi"])          # baris 3
        ws.append([dt.date(2026, 10, 21), "dua ratus", None, None])  # baris 4
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_tiket() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master Data 2026"
    for _ in range(6):
        ws.append(["LAPORAN MASTER DATA"])
    ws.append(["Ticket Date", "Ticket No", "Supplier Code", "Supplier Name",
               "Vehicle No", "Item Code", "DRC Date", "Wet", "Sample", "Reject",
               "DRC", "Dry", "Asal Getah", "Provinsi", "Kabupaten/Kota",
               "Kecamatan", "Desa", "BIN", "Location"])
    ws.append([])
    t = dt.date(2026, 10, 16)
    ws.append([t, "T900", "V5", "VENDOR LIMA", "BK 5", "SLAB2", t, 5000, 5, 10, 60,
               3000, "SUMUT, A, B, C", "SUMUT", "A", "B", "C", "BIN1", "L1"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pratinjau(klien, isi: bytes, nama: str = "stok.xlsx", pengisi="tester"):
    return klien.post(
        "/api/unggah/pratinjau",
        files={"file": (nama, isi,
                        "application/vnd.openxmlformats-officedocument"
                        ".spreadsheetml.sheet")},
        data={"pengisi": pengisi},
    )


def _hitung(Sesi, model) -> int:
    with Sesi() as s:
        return s.scalar(select(func.count()).select_from(model))


# --------------------------------------------------------------- pratinjau

def test_pratinjau_menutup_workbook_sebelum_file_dipindah(klien, monkeypatch):
    """Windows mengunci .xlsx sampai semua workbook openpyxl ditutup."""
    c, _, _ = klien
    asli = openpyxl.load_workbook
    penutup = []

    def pantau(*args, **kwargs):
        wb = asli(*args, **kwargs)
        from unittest.mock import Mock
        wb.close = Mock(wraps=wb.close)
        penutup.append(wb.close)
        return wb

    monkeypatch.setattr(openpyxl, "load_workbook", pantau)
    r = _pratinjau(c, _xlsx_stok(dengan_mismatch=False))
    assert r.status_code == 200, r.text
    assert penutup and all(tutup.call_count == 1 for tutup in penutup)


def test_pratinjau_tidak_menyimpan_apa_pun(klien):
    c, Sesi, _ = klien
    r = _pratinjau(c, _xlsx_stok())
    assert r.status_code == 200, r.text
    j = r.json()

    assert j["status"] == "pratinjau"
    assert j["token"]
    assert j["jenis"].startswith("Data berkala")
    assert j["total_jumlah"] == 1
    assert j["total_ditolak"] == 2
    assert j["jumlah_mismatch"] == 2
    assert j["jumlah"]["stok"] == 1
    assert j["ditolak"] == {"Stok_Harian": 2}
    assert j["kedaluwarsa_menit"] == 60

    # Tidak satu baris pun tersimpan.
    assert _hitung(Sesi, Stok) == 0
    assert _hitung(Sesi, RawUnggahan) == 0
    assert _hitung(Sesi, RawBaris) == 0
    assert _hitung(Sesi, PeringatanData) == 0


def test_mismatch_berisi_lokasi_dan_alasan(klien):
    c, _, _ = klien
    j = _pratinjau(c, _xlsx_stok()).json()
    mismatch = j["mismatch"]
    assert [m["baris"] for m in mismatch] == [3, 4]
    assert {m["sheet"] for m in mismatch} == {"Stok_Harian"}
    assert "Tanggal kosong" in mismatch[0]["alasan"]
    assert "tidak terbaca" in mismatch[1]["alasan"]
    assert all(m["kode"] for m in mismatch)
    assert "TIDAK akan dimasukkan" in j["pesan"]


def test_pratinjau_file_bersih_tanpa_mismatch(klien):
    c, _, _ = klien
    j = _pratinjau(c, _xlsx_stok(dengan_mismatch=False)).json()
    assert j["total_ditolak"] == 0
    assert j["jumlah_mismatch"] == 0
    assert "tidak ada baris bermasalah" in j["pesan"].lower()


def test_pratinjau_master_data_tiket(klien):
    c, Sesi, _ = klien
    j = _pratinjau(c, _xlsx_tiket(), nama="master data.xlsx").json()
    assert j["jenis"].startswith("Tiket intake")
    assert j["jumlah"]["tiket_baris"] == 1
    assert j["total_ditolak"] == 0
    assert _hitung(Sesi, RawBaris) == 0


# ------------------------------------------------------- unduh data mismatch

def test_unduh_excel_mismatch(klien):
    c, _, _ = klien
    token = _pratinjau(c, _xlsx_stok()).json()["token"]
    r = c.get(f"/api/unggah/pratinjau/{token}/mismatch.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    assert "Ringkasan" in wb.sheetnames
    ringkas = wb["Ringkasan"]
    assert [c.value for c in ringkas[1]][:4] == [
        "SHEET", "BARIS", "KODE MASALAH", "ALASAN (yang harus diperbaiki)"]
    assert ringkas["A2"].value == "Stok_Harian"
    assert ringkas["B2"].value == 3
    assert "Tanggal kosong" in ringkas["D2"].value
    assert ringkas.max_row == 3  # 2 baris mismatch

    # Lembar per sheet: header asli dipertahankan + 3 kolom tambahan
    lembar = wb["Stok_Harian"]
    judul = [c.value for c in lembar[1]]
    assert judul[:4] == ["Tanggal", "Stok_Kg", "Item_Code", "Catatan"]
    assert judul[-3:] == ["KODE_MASALAH", "ALASAN_DITOLAK",
                          "BARIS_DI_FILE_ASLI"]
    assert lembar["B2"].value == 100            # nilai asli ikut tersalin
    assert lembar["D2"].value == "tanggal lupa diisi"
    assert lembar["G2"].value == 3              # nomor baris asli
    assert "Tanggal kosong" in lembar["F2"].value
    assert lembar["B3"].value == "dua ratus"
    assert "Cara memperbaiki" in wb.sheetnames


def test_unduh_mismatch_tanpa_baris_bermasalah(klien):
    c, _, _ = klien
    token = _pratinjau(c, _xlsx_stok(dengan_mismatch=False)).json()["token"]
    r = c.get(f"/api/unggah/pratinjau/{token}/mismatch.xlsx")
    assert r.status_code == 404


# --------------------------------------------------------------- konfirmasi

def test_file_preview_terkunci_memberi_pesan_terbaca_dan_tidak_menghapus(klien, monkeypatch):
    c, Sesi, folder = klien
    token = _pratinjau(c, _xlsx_stok(dengan_mismatch=False)).json()["token"]

    def terkunci(*args, **kwargs):
        raise PermissionError("file sedang digunakan")

    monkeypatch.setattr("app.api_unggah.shutil.move", terkunci)
    r = c.post(f"/api/unggah/pratinjau/{token}/konfirmasi")
    assert r.status_code == 423
    assert "masih sedang dipakai" in r.json()["detail"]
    assert _hitung(Sesi, Stok) == 0
    assert list(folder.rglob("*.xlsx"))


def test_konfirmasi_baru_menyimpan_ke_database(klien):
    c, Sesi, folder = klien
    token = _pratinjau(c, _xlsx_stok()).json()["token"]
    assert _hitung(Sesi, Stok) == 0

    r = c.post(f"/api/unggah/pratinjau/{token}/konfirmasi")
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["status"] == "selesai"
    assert j["jumlah"]["stok"] == 1
    assert j["ditolak"] == {"Stok_Harian": 2}
    assert j["jumlah_mismatch"] == 2
    assert j["unggahan_id"] > 0
    assert "masuk database" in j["pesan"]

    with Sesi() as s:
        baris = s.scalars(select(Stok)).all()
        assert len(baris) == 1
        assert baris[0].tanggal == dt.date(2026, 10, 20)
        assert str(baris[0].stok_kg) == "500000.000"
        unggahan = s.scalars(select(RawUnggahan)).all()
        assert len(unggahan) == 1
        assert unggahan[0].pengunggah == "tester"

    # File asli tersimpan permanen, file sementara pratinjau sudah hilang.
    permanen = [p.name for p in folder.iterdir() if p.is_file()]
    assert len(permanen) == 1
    assert permanen[0].endswith("stok.xlsx")


def test_token_sekali_pakai(klien):
    c, _, _ = klien
    token = _pratinjau(c, _xlsx_stok()).json()["token"]
    assert c.post(f"/api/unggah/pratinjau/{token}/konfirmasi").status_code == 200
    # Konfirmasi kedua tidak boleh menggandakan data.
    assert c.post(f"/api/unggah/pratinjau/{token}/konfirmasi").status_code == 409
    # Daftar mismatch baris yang ditolak masih bisa diunduh untuk diperbaiki.
    assert c.get(f"/api/unggah/pratinjau/{token}/mismatch.xlsx").status_code == 200


def test_konfirmasi_master_data_tiket(klien):
    c, Sesi, _ = klien
    token = _pratinjau(c, _xlsx_tiket(), nama="master data.xlsx").json()["token"]
    j = c.post(f"/api/unggah/pratinjau/{token}/konfirmasi").json()
    assert j["jumlah"]["tiket_baris"] == 1
    assert _hitung(Sesi, RawBaris) > 0


# ------------------------------------------------------------------ batalkan

def test_batalkan_tidak_menyimpan_apa_pun(klien):
    c, Sesi, folder = klien
    token = _pratinjau(c, _xlsx_stok()).json()["token"]
    r = c.delete(f"/api/unggah/pratinjau/{token}")
    assert r.status_code == 200
    assert r.json()["status"] == "dibatalkan"

    assert _hitung(Sesi, Stok) == 0
    assert [p for p in folder.rglob("*.xlsx")] == []
    assert c.post(f"/api/unggah/pratinjau/{token}/konfirmasi").status_code == 410


def test_token_asing_ditolak(klien):
    c, _, _ = klien
    assert c.post("/api/unggah/pratinjau/ngawur/konfirmasi").status_code == 410
    assert c.get("/api/unggah/pratinjau/ngawur/mismatch.xlsx").status_code == 410


# ------------------------------------------------------------ berkas salah

def test_bukan_xlsx_ditolak(klien):
    c, _, _ = klien
    r = _pratinjau(c, b"bukan excel", nama="data.csv")
    assert r.status_code == 422
    assert ".xlsx" in r.json()["detail"]


def test_xlsx_rusak_ditolak(klien):
    c, _, folder = klien
    r = _pratinjau(c, b"bukan excel sama sekali", nama="rusak.xlsx")
    assert r.status_code == 422
    assert "tidak bisa dibuka" in r.json()["detail"]
    assert [p for p in folder.rglob("*.xlsx")] == []


def test_file_kosong_ditolak(klien):
    c, _, _ = klien
    r = _pratinjau(c, b"", nama="kosong.xlsx")
    assert r.status_code == 422
    assert "kosong" in r.json()["detail"].lower()
