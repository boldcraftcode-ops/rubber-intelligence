"""Tes analisis deskriptif & komposisi penerimaan (share, Pareto/ABC, EUDR).

Permintaan pemilik: satu analisa lagi yang **langsung bisa dipakai** —
share per provinsi (ketergantungan wilayah), share EUDR vs biasa (dipantau per
bulan untuk target kepatuhan), dan Pareto/ABC (siapa yang menyumbang 80%).

Yang dijaga:
- porsi dihitung dari total periode terpilih, jumlahnya 100%;
- kumulatif naik monoton dan kelas ABC mengikuti ambang 80% / 95% (kategori
  yang MELEWATI ambang tetap masuk kelas sebelumnya, tidak terbelah);
- jumlah kategori "diperlukan untuk 80%" menghitung kategori yang membuat
  kumulatif mencapai 80%, bukan yang kumulatifnya ≤80%;
- share EUDR dihitung dari grade ber-flag EUDR/TTSR, total dan per bulan;
- wilayah memakai kunci ter-normalisasi (UPPER/TRIM) seperti kartu Penerimaan.
"""
import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import buat_app
from app.models import Base, Grade, TiketBaris, Vendor


def _engine():
    eng = create_engine("sqlite://", future=True,
                        connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    Base.metadata.create_all(eng)
    return eng


def _tiket(session, no, tanggal, kg, item="SIR20", vendor="V001",
           provinsi="Sumatera Utara", kabupaten="Deli Serdang"):
    session.add(TiketBaris(
        ticket_no=no, line_seq=1, ticket_date=tanggal, kode_vendor=vendor,
        item_code=item, dry_kg=kg, wet_kg=kg / 0.62,
        provinsi=provinsi, kabupaten=kabupaten, tiket_nol=False))


def _klien(monkeypatch, isi):
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("AUTH_DISABLED", "true")
    eng = _engine()
    with sessionmaker(bind=eng, future=True)() as s:
        s.add_all([Vendor(kode="V001", nama_baku="PT Alpha"),
                   Vendor(kode="V002", nama_baku="PT Beta"),
                   Grade(item_code="SIR20", kelompok="SIR", eudr=False),
                   Grade(item_code="TTSR20", kelompok="TSR", eudr=True),
                   Grade(item_code="SIR10", kelompok="SIR", eudr=False)])
        isi(s)
        s.commit()
    return TestClient(buat_app(eng))


def _komp(klien, **params):
    r = klien.get("/api/kartu/penerimaan/komposisi", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def _isi_tiga_provinsi(s):
    """Sumut 50%, Riau 25%, Aceh 25% dari total 1.000.000 kg."""
    _tiket(s, "P1", dt.date(2026, 1, 5), 500_000, provinsi="Sumatera Utara")
    _tiket(s, "P2", dt.date(2026, 1, 6), 250_000, provinsi="Riau",
           kabupaten="Kampar")
    _tiket(s, "P3", dt.date(2026, 1, 7), 250_000, provinsi="Aceh",
           kabupaten="Aceh Timur")


# ================================================================== share/porsi
def test_share_per_provinsi(monkeypatch):
    d = _komp(_klien(monkeypatch, _isi_tiga_provinsi))
    assert d["total_kg"] == pytest.approx(1_000_000.0)
    assert d["total_ton"] == pytest.approx(1000.0)
    assert [b["label"] for b in d["baris"]] == ["Sumatera Utara", "Riau", "Aceh"]
    porsi = [b["porsi_pct"] for b in d["baris"]]
    assert porsi == [50.0, 25.0, 25.0]
    assert sum(porsi) == pytest.approx(100.0)
    assert d["baris"][0]["kumulatif_pct"] == 50.0
    assert d["baris"][-1]["kumulatif_pct"] == pytest.approx(100.0)


def test_ketergantungan_satu_wilayah(monkeypatch):
    d = _komp(_klien(monkeypatch, _isi_tiga_provinsi))
    k = d["konsentrasi"]
    assert k["top1_label"] == "Sumatera Utara"
    assert k["top1_porsi_pct"] == 50.0
    assert k["top3_porsi_pct"] == pytest.approx(100.0)
    # 0,5² + 0,25² + 0,25² = 0,375
    assert k["hhi"] == pytest.approx(0.375, abs=1e-4)
    assert "terkonsentrasi" in k["label"]


def test_hhi_rendah_bila_menyebar(monkeypatch):
    def isi(s):
        for i in range(10):
            _tiket(s, f"M{i}", dt.date(2026, 1, 5 + i), 100_000,
                   provinsi=f"Provinsi {i}")
    k = _komp(_klien(monkeypatch, isi))["konsentrasi"]
    assert k["hhi"] == pytest.approx(0.1, abs=1e-4)
    assert k["label"] == "rendah (pasokan menyebar)"


def test_provinsi_dobel_digabung_dalam_share(monkeypatch):
    def isi(s):
        _tiket(s, "D1", dt.date(2026, 1, 5), 300_000, provinsi="Sumatera Utara")
        _tiket(s, "D2", dt.date(2026, 1, 6), 100_000, provinsi="SUMATERA UTARA")
        _tiket(s, "D3", dt.date(2026, 1, 7), 100_000, provinsi="Riau")
    d = _komp(_klien(monkeypatch, isi))
    assert len(d["baris"]) == 2
    su = d["baris"][0]
    assert su["label"] == "Sumatera Utara"
    assert su["kg"] == pytest.approx(400_000.0)
    assert su["porsi_pct"] == 80.0
    assert any("digabung" in p for p in d["peringatan"])


# ================================================================= Pareto / ABC
def test_pareto_ambang_80_dan_kelas(monkeypatch):
    """60/20/15/5 -> A: 60+20=80 (2 kategori), B: +15=95 (1), C: 5 (1)."""
    def isi(s):
        for i, kg in enumerate([600_000, 200_000, 150_000, 50_000]):
            _tiket(s, f"A{i}", dt.date(2026, 1, 5 + i), kg,
                   provinsi=f"Wilayah {i}")
    d = _komp(_klien(monkeypatch, isi))
    assert [b["kelas"] for b in d["baris"]] == ["A", "A", "B", "C"]
    assert d["pareto"]["jumlah_kategori_80"] == 2   # 60+20 tepat mencapai 80%
    assert [b["kumulatif_pct"] for b in d["baris"]] == [60.0, 80.0, 95.0, 100.0]
    kelas = d["pareto"]["kelas"]
    assert kelas["A"]["jumlah"] == 2
    assert kelas["A"]["porsi_pct"] == 80.0
    assert kelas["B"]["jumlah"] == 1
    assert kelas["C"]["jumlah"] == 1
    assert "80% volume" in d["pareto"]["catatan"]


def test_jumlah_kategori_80_adalah_yang_diperlukan(monkeypatch):
    """50+25+25: 2 kategori baru 75%, jadi butuh 3 untuk mencapai 80%."""
    d = _komp(_klien(monkeypatch, _isi_tiga_provinsi))
    assert d["pareto"]["jumlah_kategori_80"] == 3
    assert d["pareto"]["kelas"]["A"]["jumlah"] == 3


def test_satu_kategori_dominan(monkeypatch):
    def isi(s):
        _tiket(s, "S1", dt.date(2026, 1, 5), 900_000, provinsi="Sumatera Utara")
        _tiket(s, "S2", dt.date(2026, 1, 6), 100_000, provinsi="Riau")
    d = _komp(_klien(monkeypatch, isi))
    assert d["pareto"]["jumlah_kategori_80"] == 1
    assert d["baris"][0]["kelas"] == "A"
    # 90% sudah lewat 80%, jadi kategori kedua masuk B (80–95%), bukan C
    assert d["baris"][1]["kelas"] == "B"


def test_pareto_vendor_dan_grade(monkeypatch):
    def isi(s):
        _tiket(s, "V1", dt.date(2026, 1, 5), 700_000, vendor="V001", item="SIR20")
        _tiket(s, "V2", dt.date(2026, 1, 6), 200_000, vendor="V002", item="SIR10")
        _tiket(s, "V3", dt.date(2026, 1, 7), 100_000, vendor="V002", item="TTSR20")
    klien = _klien(monkeypatch, isi)
    v = _komp(klien, dimensi="vendor")
    assert v["baris"][0]["label"] == "V001 — PT Alpha"
    assert v["baris"][0]["porsi_pct"] == 70.0
    assert v["pareto"]["jumlah_kategori_80"] == 2
    g = _komp(klien, dimensi="grade")
    assert {b["kunci"] for b in g["baris"]} == {"SIR20", "SIR10", "TTSR20"}


# ================================================================ EUDR vs biasa
def test_share_eudr_total(monkeypatch):
    def isi(s):
        _tiket(s, "E1", dt.date(2026, 1, 5), 800_000, item="SIR20")
        _tiket(s, "E2", dt.date(2026, 1, 6), 200_000, item="TTSR20")
    d = _komp(_klien(monkeypatch, isi))
    assert d["eudr"]["total_kg"] == pytest.approx(200_000.0)
    assert d["eudr"]["porsi_pct"] == pytest.approx(20.0)
    assert d["eudr"]["biasa_kg"] == pytest.approx(800_000.0)
    assert d["eudr"]["biasa_porsi_pct"] == pytest.approx(80.0)


def test_share_eudr_per_bulan_naik(monkeypatch):
    def isi(s):
        # Jan 10%, Feb 30%, Mar 50% EUDR
        for i, (bulan, total, eudr) in enumerate([(1, 100_000, 10_000),
                                                  (2, 100_000, 30_000),
                                                  (3, 100_000, 50_000)]):
            _tiket(s, f"B{i}a", dt.date(2026, bulan, 5), total - eudr, item="SIR20")
            _tiket(s, f"B{i}b", dt.date(2026, bulan, 6), eudr, item="TTSR20")
    d = _komp(_klien(monkeypatch, isi))
    bulan = d["eudr"]["per_bulan"]
    assert [m["label"] for m in bulan] == ["Jan-26", "Feb-26", "Mar-26"]
    assert [m["eudr_pct"] for m in bulan] == [10.0, 30.0, 50.0]
    assert d["eudr"]["porsi_pct"] == pytest.approx(30.0)


def test_tanpa_grade_eudr_ada_peringatan(monkeypatch):
    d = _komp(_klien(monkeypatch, _isi_tiga_provinsi))
    assert d["eudr"]["porsi_pct"] == 0.0
    assert any("EUDR" in p for p in d["peringatan"])


def test_eudr_per_kategori(monkeypatch):
    def isi(s):
        _tiket(s, "K1", dt.date(2026, 1, 5), 500_000, provinsi="Sumatera Utara",
               item="SIR20")
        _tiket(s, "K2", dt.date(2026, 1, 6), 500_000, provinsi="Sumatera Utara",
               item="TTSR20")
    d = _komp(_klien(monkeypatch, isi))
    assert d["baris"][0]["eudr_kg"] == pytest.approx(500_000.0)
    assert d["baris"][0]["eudr_pct"] == pytest.approx(50.0)


# ==================================================================== lain-lain
def test_filter_periode_dan_mode_eudr(monkeypatch):
    def isi(s):
        _tiket(s, "T1", dt.date(2026, 1, 5), 100_000)
        _tiket(s, "T2", dt.date(2026, 3, 5), 300_000)
        _tiket(s, "T3", dt.date(2026, 3, 6), 100_000, item="TTSR20")
    klien = _klien(monkeypatch, isi)
    d = _komp(klien, dari="2026-03", sampai="2026-03")
    assert d["total_kg"] == pytest.approx(400_000.0)
    hanya_eudr = _komp(klien, eudr="true")
    assert hanya_eudr["total_kg"] == pytest.approx(100_000.0)
    assert hanya_eudr["eudr"]["porsi_pct"] == pytest.approx(100.0)


def test_dimensi_tidak_dikenal_422(monkeypatch):
    klien = _klien(monkeypatch, _isi_tiga_provinsi)
    assert klien.get("/api/kartu/penerimaan/komposisi",
                     params={"dimensi": "bintang"}).status_code == 422


def test_periode_terbalik_422(monkeypatch):
    klien = _klien(monkeypatch, _isi_tiga_provinsi)
    r = klien.get("/api/kartu/penerimaan/komposisi",
                  params={"dari": "2026-05", "sampai": "2026-01"})
    assert r.status_code == 422
    assert "terbalik" in r.json()["detail"]


def test_tanpa_data_tidak_galat(monkeypatch):
    d = _komp(_klien(monkeypatch, lambda s: None))
    assert d["total_kg"] == 0
    assert d["baris"] == []
    assert d["eudr"]["porsi_pct"] is None
    assert any("Tidak ada data" in p for p in d["peringatan"])


def test_ekspor_xlsx(monkeypatch):
    klien = _klien(monkeypatch, _isi_tiga_provinsi)
    r = klien.get("/api/kartu/penerimaan/komposisi", params={"format": "xlsx"})
    assert r.status_code == 200
    assert "komposisi_penerimaan_provinsi.xlsx" in r.headers["content-disposition"]
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["Provinsi", "Kg", "Porsi %", "Kumulatif %",
                                        "Kelas ABC", "EUDR kg", "EUDR %",
                                        "Jumlah tiket"]
    assert ws.max_row == 4       # header + 3 provinsi


def test_health_menyebut_fitur_komposisi(monkeypatch):
    """Layar 'beda versi' memakai /api/health untuk mendeteksi backend lama."""
    fitur = _klien(monkeypatch, _isi_tiga_provinsi).get("/api/health").json()["fitur"]
    assert "penerimaan-komposisi" in fitur
    assert "penerimaan-tren" in fitur
