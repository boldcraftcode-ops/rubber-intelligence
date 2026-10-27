"""Tes analisis tren & musiman pada kartu Penerimaan karet, plus normalisasi
nama wilayah.

Latar (permintaan pemilik):
- sematkan analisis tren & musiman di kartu "Penerimaan karet (kg) — biasa vs
  EUDR": moving average 3 bulan, dekomposisi tren-musiman (STL), MoM & YoY;
- duplikasi nama provinsi ("Sumatera Utara" vs "SUMATERA UTARA", "Riau" vs
  "RIAU") hampir pasti masalah standardisasi huruf besar/kecil;
- pola datar tidak wajar (Jan-Jun nyaris konstan) harus ditandai sebelum
  angka dipakai untuk forecasting.

Yang dijaga:
- ejaan wilayah yang berbeda tetap SATU baris, dan penggabungannya dilaporkan;
- filter/drill wilayah tidak peka huruf besar/kecil;
- MA3, MoM, YoY dihitung benar, dan kosong (bukan 0%) saat pembanding tidak ada;
- dekomposisi menolak jalan di bawah 24 bulan, dan memulihkan pola musiman
  yang ditanam di data uji saat data cukup;
- nol, deret nol di tengah, pola datar, dan bulan terakhir belum lengkap
  dilaporkan sebagai peringatan.
"""
import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import buat_app
from app.models import Base, Grade, TiketBaris, Vendor


def _engine():
    eng = create_engine("sqlite://", future=True,
                        connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    Base.metadata.create_all(eng)
    return eng


def _tiket(session, no, tanggal, kg, item="SIR20", vendor="V001",
           provinsi="Sumatera Utara", kabupaten="Deli Serdang"):
    session.add(TiketBaris(
        ticket_no=no, line_seq=1, ticket_date=tanggal, kode_vendor=vendor,
        item_code=item, dry_kg=kg, wet_kg=kg / 0.62,
        provinsi=provinsi, kabupaten=kabupaten, tiket_nol=False))


def _klien(monkeypatch, isi):
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("AUTH_DISABLED", "true")
    eng = _engine()
    with sessionmaker(bind=eng, future=True)() as s:
        s.add_all([Vendor(kode="V001", nama_baku="PT Alpha"),
                   Vendor(kode="V002", nama_baku="PT Beta"),
                   Grade(item_code="SIR20", kelompok="SIR", eudr=False),
                   Grade(item_code="TTSR20", kelompok="TSR", eudr=True)])
        isi(s)
        s.commit()
    return TestClient(buat_app(eng))


def _tren(klien, **params):
    r = klien.get("/api/kartu/penerimaan/tren", params=params)
    assert r.status_code == 200, r.text
    return r.json()


# ======================================================== normalisasi wilayah
def _isi_dua_ejaan(s):
    # provinsi sama, ejaan berbeda; RIAU juga dobel
    _tiket(s, "A1", dt.date(2026, 10, 1), 5_000, provinsi="Sumatera Utara")
    _tiket(s, "A2", dt.date(2026, 10, 2), 7_100, provinsi="SUMATERA UTARA")
    _tiket(s, "A3", dt.date(2026, 10, 3), 4_680, provinsi="RIAU")
    _tiket(s, "A4", dt.date(2026, 10, 4), 2_000, provinsi="Riau")
    _tiket(s, "A5", dt.date(2026, 10, 5), 3_000, provinsi="  sumatera   utara ")


def test_provinsi_dobel_jadi_satu_baris(monkeypatch):
    klien = _klien(monkeypatch, _isi_dua_ejaan)
    d = klien.get("/api/kartu/penerimaan", params={"dimensi": "provinsi"}).json()
    kunci = [b["kunci"] for b in d["baris"]]
    assert kunci == ["SUMATERA UTARA", "RIAU"]      # hanya dua, bukan empat
    su = next(b for b in d["baris"] if b["kunci"] == "SUMATERA UTARA")
    assert su["kg_total"] == pytest.approx(15_100.0)   # 5.000 + 7.100 + 3.000
    assert su["label"] == "Sumatera Utara"             # ejaan paling wajar
    assert su["jumlah_tiket"] == 3


def test_laporan_penggabungan_wilayah(monkeypatch):
    klien = _klien(monkeypatch, _isi_dua_ejaan)
    d = klien.get("/api/kartu/penerimaan", params={"dimensi": "provinsi"}).json()
    lap = d["wilayah"]
    assert lap["dinormalisasi"] is True
    digabung = {g["kunci"]: g for g in lap["digabung"]}
    assert set(digabung) == {"SUMATERA UTARA", "RIAU"}
    ejaan_su = {v["ejaan"] for v in digabung["SUMATERA UTARA"]["varian"]}
    assert "Sumatera Utara" in ejaan_su and "SUMATERA UTARA" in ejaan_su
    assert "digabung" in lap["catatan"].lower()


def test_normalisasi_juga_di_mode_bulan_dan_tonase(monkeypatch):
    klien = _klien(monkeypatch, _isi_dua_ejaan)
    d = klien.get("/api/kartu/penerimaan",
                  params={"dimensi": "provinsi", "per": "bulan"}).json()
    assert [b["kunci"] for b in d["baris"]] == ["SUMATERA UTARA", "RIAU"]
    assert d["baris"][0]["label"] == "Sumatera Utara"
    assert d["wilayah"]["digabung"]

    t = klien.get("/api/kartu/tonase", params={"dimensi": "provinsi"}).json()
    assert [b["kunci"] for b in t["baris"]] == ["SUMATERA UTARA", "RIAU"]
    assert t["baris"][0]["label"] == "Sumatera Utara"


def test_filter_wilayah_tidak_peka_huruf_besar(monkeypatch):
    klien = _klien(monkeypatch, _isi_dua_ejaan)
    for nilai in ("Sumatera Utara", "SUMATERA UTARA", "sumatera utara"):
        d = klien.get("/api/kartu/penerimaan",
                      params={"dimensi": "kabupaten", "provinsi": nilai}).json()
        assert d["total_kg"] == pytest.approx(15_100.0), nilai


# ================================================================== MA3 / MoM / YoY
def _isi_lima_bulan(s):
    for i, kg in enumerate([100_000, 200_000, 300_000, 400_000, 500_000]):
        _tiket(s, f"M{i}", dt.date(2026, i + 1, 10), kg)


def test_ma3_dan_mom(monkeypatch):
    klien = _klien(monkeypatch, _isi_lima_bulan)
    d = _tren(klien)
    assert d["total_kg"] == [100_000.0, 200_000.0, 300_000.0, 400_000.0, 500_000.0]
    assert d["ma3"] == [None, None, 200_000.0, 300_000.0, 400_000.0]
    assert d["mom_pct"] == [None, 100.0, 50.0, 33.33, 25.0]


def test_yoy_kosong_bila_tahun_lalu_belum_ada(monkeypatch):
    klien = _klien(monkeypatch, _isi_lima_bulan)
    assert _tren(klien)["yoy_pct"] == [None] * 5


def test_yoy_dihitung_bila_ada_data_tahun_lalu(monkeypatch):
    def isi(s):
        _tiket(s, "Y1", dt.date(2025, 1, 10), 100_000)
        _tiket(s, "Y2", dt.date(2025, 2, 10), 100_000)
        _tiket(s, "Y3", dt.date(2026, 1, 10), 150_000)
        _tiket(s, "Y4", dt.date(2026, 2, 10), 50_000)
    klien = _klien(monkeypatch, isi)
    d = _tren(klien)
    yoy = {b["ym"]: v for b, v in zip(d["bulan"], d["yoy_pct"])}
    assert yoy["2026-01"] == pytest.approx(50.0)
    assert yoy["2026-02"] == pytest.approx(-50.0)
    assert yoy["2025-01"] is None


def test_mom_kosong_saat_pembanding_nol(monkeypatch):
    def isi(s):
        _tiket(s, "Z1", dt.date(2026, 1, 10), 100_000)
        _tiket(s, "Z3", dt.date(2026, 3, 10), 80_000)
    klien = _klien(monkeypatch, isi)
    d = _tren(klien)
    assert d["total_kg"] == [100_000.0, 0.0, 80_000.0]
    assert d["mom_pct"][2] is None          # Februari nol -> bukan -100%


# ============================================================== dekomposisi
def _isi_tiga_tahun(s, bulan_mulai=1, jumlah=36):
    """Tren naik + musiman buatan: Feb-Apr rendah, Jul-Agu tinggi."""
    for i in range(jumlah):
        th = 2023 + (bulan_mulai - 1 + i) // 12
        bl = (bulan_mulai - 1 + i) % 12 + 1
        nilai = 100_000 + 2_000 * i
        if bl in (2, 3, 4):
            nilai -= 8_000
        elif bl in (7, 8):
            nilai += 3_000
        _tiket(s, f"T{i:03d}", dt.date(th, bl, 10), nilai)


def test_dekomposisi_butuh_minimal_24_bulan(monkeypatch):
    klien = _klien(monkeypatch, lambda s: _isi_tiga_tahun(s, jumlah=18))
    d = _tren(klien)["dekomposisi"]
    assert d["tersedia"] is False
    assert d["min_bulan"] == 24
    assert d["jumlah_bulan"] == 18
    assert d["kurang_bulan"] == 6
    assert "24 bulan" in d["pesan"]


def test_dekomposisi_jalan_dengan_36_bulan(monkeypatch):
    klien = _klien(monkeypatch, lambda s: _isi_tiga_tahun(s, jumlah=36))
    d = _tren(klien)["dekomposisi"]
    assert d["tersedia"] is True
    assert d["jumlah_bulan"] == 36
    assert d["metode"].lower().startswith("dekomposisi aditif")
    assert len(d["tren"]) == 36 and len(d["musiman"]) == 36


def test_dekomposisi_memulihkan_bulan_musiman(monkeypatch):
    klien = _klien(monkeypatch, lambda s: _isi_tiga_tahun(s, jumlah=36))
    d = _tren(klien)["dekomposisi"]
    indeks = {b["bulan"]: b["kg"] for b in d["indeks_musiman"]}
    # Feb-Apr paling negatif, Jul-Agu paling positif — sesuai pola yang ditanam
    assert all(indeks[m] < 0 for m in (2, 3, 4))
    assert all(indeks[m] > 0 for m in (7, 8))
    assert min(indeks, key=indeks.get) in (2, 3, 4)
    assert d["kekuatan_musiman_pct"] > 0


def test_dekomposisi_tren_naik_mengikuti_data(monkeypatch):
    klien = _klien(monkeypatch, lambda s: _isi_tiga_tahun(s, jumlah=36))
    d = _tren(klien)["dekomposisi"]
    tren = [x for x in d["tren"] if x is not None]
    assert len(tren) >= 20
    assert tren[0] < tren[-1]         # tren dasar memang naik


# ================================================================ validasi data
def test_bulan_nol_dan_rentetan_nol_di_tengah(monkeypatch):
    def isi(s):
        _tiket(s, "N1", dt.date(2026, 1, 10), 100_000)
        _tiket(s, "N5", dt.date(2026, 5, 10), 90_000)
    klien = _klien(monkeypatch, isi)
    v = _tren(klien)["validasi"]
    assert v["bulan_nol"] == ["2026-02", "2026-03", "2026-04"]
    assert v["bulan_nol_di_tengah"] == ["2026-02", "2026-03", "2026-04"]
    assert any("belum diunggah" in p for p in v["peringatan"])


def test_pola_datar_ditandai(monkeypatch):
    def isi(s):
        for i, kg in enumerate([190_000, 191_000, 189_000, 192_000, 190_500, 191_000]):
            _tiket(s, f"F{i}", dt.date(2026, i + 1, 10), kg)
    klien = _klien(monkeypatch, isi)
    v = _tren(klien)["validasi"]
    assert len(v["pola_datar"]) == 1
    d = v["pola_datar"][0]
    assert d["jumlah_bulan"] == 6
    assert d["mulai"] == "2026-01" and d["selesai"] == "2026-06"
    assert d["cv_pct"] < 5
    assert any("Pola datar tidak wajar" in p for p in v["peringatan"])


def test_pola_wajar_tidak_ditandai(monkeypatch):
    def isi(s):
        for i, kg in enumerate([120_000, 180_000, 90_000, 210_000]):
            _tiket(s, f"W{i}", dt.date(2026, i + 1, 10), kg)
    klien = _klien(monkeypatch, isi)
    assert _tren(klien)["validasi"]["pola_datar"] == []


def test_bulan_terakhir_belum_lengkap(monkeypatch):
    def isi(s):
        _tiket(s, "L1", dt.date(2026, 1, 31), 100_000)
        _tiket(s, "L2", dt.date(2026, 2, 10), 50_000)   # berhenti tgl 10
    klien = _klien(monkeypatch, isi)
    v = _tren(klien)["validasi"]
    assert v["bulan_terakhir_lengkap"] is False
    assert v["tanggal_terakhir"] == "2026-02-10"
    assert any("belum lengkap" in p for p in v["peringatan"])


def test_bulan_terakhir_lengkap(monkeypatch):
    def isi(s):
        _tiket(s, "K1", dt.date(2026, 1, 31), 100_000)
    klien = _klien(monkeypatch, isi)
    assert _tren(klien)["validasi"]["bulan_terakhir_lengkap"] is True


# ================================================================== lain-lain
def test_periode_memotong_tampilan_tetapi_ma3_tetap_ada(monkeypatch):
    """Bulan pertama dalam periode tetap punya MA3 dari bulan sebelumnya."""
    klien = _klien(monkeypatch, lambda s: _isi_tiga_tahun(s, jumlah=36))
    d = _tren(klien, dari="2025-07", sampai="2025-12")
    assert d["jumlah_bulan"] == 6
    assert d["bulan"][0]["ym"] == "2025-07"
    assert d["bulan"][-1]["ym"] == "2025-12"
    assert d["ma3"][0] is not None          # dihitung dari deret penuh


def test_periode_terbalik_422(monkeypatch):
    klien = _klien(monkeypatch, _isi_lima_bulan)
    r = klien.get("/api/kartu/penerimaan/tren",
                  params={"dari": "2026-05", "sampai": "2026-01"})
    assert r.status_code == 422
    assert "terbalik" in r.json()["detail"]


def test_pemisahan_biasa_dan_eudr(monkeypatch):
    def isi(s):
        _tiket(s, "E1", dt.date(2026, 1, 10), 60_000, item="SIR20")
        _tiket(s, "E2", dt.date(2026, 1, 11), 40_000, item="TTSR20")
    klien = _klien(monkeypatch, isi)
    d = _tren(klien)
    assert d["total_kg"] == [100_000.0]
    assert d["biasa_kg"] == [60_000.0]
    assert d["eudr_kg"] == [40_000.0]
    hanya_eudr = _tren(klien, eudr="true")
    assert hanya_eudr["total_kg"] == [40_000.0]


def test_tanpa_data_tidak_galat(monkeypatch):
    klien = _klien(monkeypatch, lambda s: None)
    d = klien.get("/api/kartu/penerimaan/tren").json()
    assert d["status"] == "kosong"
    assert d["bulan"] == []
    assert d["dekomposisi"]["tersedia"] is False


def test_ekspor_xlsx(monkeypatch):
    klien = _klien(monkeypatch, _isi_lima_bulan)
    r = klien.get("/api/kartu/penerimaan/tren", params={"format": "xlsx"})
    assert r.status_code == 200
    assert "penerimaan_tren_musiman.xlsx" in r.headers["content-disposition"]
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["Bulan", "Total kg", "Biasa kg",
                                        "EUDR kg", "MA3 kg", "MoM %", "YoY %",
                                        "Jumlah tiket"]
    assert ws.max_row == 6                  # header + 5 bulan
