"""Tes menu Master Vendor: CRUD, template bulk import, dan impor massal.

Permintaan pemilik: "buat master vendor dimana pada menu itu bisa melakukan
CRUD, sediakan juga template untuk bulk import".

Yang dijaga:
- kode vendor = kunci: disimpan HURUF BESAR, tidak boleh dobel, tidak boleh
  berisi karakter aneh (karena harus sama persis dengan kode di tiket);
- vendor yang sudah dipakai tiket/pembelian TIDAK terhapus diam-diam: 409
  dengan pesan yang menyebut jumlah pemakaian, dan hapus paksa tetap dicatat;
- template .xlsx benar-benar bisa dibuka, punya sheet `Master_Vendor`, dan
  hasil unduhannya bisa langsung diunggah kembali (bolak-balik);
- impor massal bersifat upsert: kode lama diperbarui, kode baru ditambah,
  baris rusak dilaporkan tanpa membatalkan baris lain;
- impor tidak pernah membuat vendor tanpa nama.
"""
import datetime as dt
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import buat_app
from app.models import Base, Pembelian, TiketBaris, Vendor


def _engine():
    eng = create_engine("sqlite://", future=True,
                        connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    Base.metadata.create_all(eng)
    return eng


@pytest.fixture()
def klien(monkeypatch):
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("AUTH_DISABLED", "true")
    return TestClient(buat_app(_engine()))


@pytest.fixture()
def klien_terpakai(monkeypatch):
    """Klien dengan vendor V001 yang sudah dipakai tiket & pembelian."""
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("AUTH_DISABLED", "true")
    eng = _engine()
    with sessionmaker(bind=eng, future=True)() as s:
        s.add(Vendor(kode="V001", nama_baku="PT Alpha", kategori="nasional",
                     status="aktif"))
        s.add(TiketBaris(ticket_no="T001", line_seq=1,
                         ticket_date=dt.date(2026, 1, 5), kode_vendor="V001",
                         item_code="SIR20", dry_kg=12_000, wet_kg=20_000,
                         tiket_nol=False))
        s.add(Pembelian(tanggal=dt.date(2026, 1, 5), item_code="SIR20",
                        kode_vendor="V001", ticket_no="T001",
                        qty_kering_kg=12_000, amount_rp=100_000))
        s.commit()
    return TestClient(buat_app(eng))


def _xlsx(sheet: str = "Master_Vendor", baris: list[list] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Kode_Vendor", "Nama_Baku", "Varian_Nama", "Kategori", "Status"])
    for b in baris or []:
        ws.append(b)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _unggah(klien, isi: bytes, nama: str = "vendor.xlsx"):
    return klien.post("/api/vendor/impor",
                      files={"file": (nama, isi,
                                      "application/vnd.openxmlformats-officedocument"
                                      ".spreadsheetml.sheet")},
                      data={"pengisi": "budi"})


# ============================================================ CREATE & READ
def test_tambah_vendor(klien):
    r = klien.post("/api/vendor", json={
        "kode": "v002", "nama_baku": "  CV Beta  ", "kategori": "lokal",
        "status": "aktif", "pengisi": "budi"})
    assert r.status_code == 200, r.text
    v = r.json()["vendor"]
    assert v["kode"] == "V002"             # kode dibakukan HURUF BESAR
    assert v["nama_baku"] == "CV Beta"     # spasi tepi dibuang
    assert v["terpakai"] is False


def test_kode_dobel_ditolak(klien):
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Alpha"})
    r = klien.post("/api/vendor", json={"kode": "v001", "nama_baku": "PT Lain"})
    assert r.status_code == 409
    assert "sudah ada" in r.json()["detail"].lower()


def test_kode_tidak_sah_ditolak(klien):
    for kode in ["V 001", "V#001", "", "X" * 21]:
        r = klien.post("/api/vendor", json={"kode": kode, "nama_baku": "PT X"})
        assert r.status_code == 422, kode


def test_status_tidak_dikenal_ditolak(klien):
    r = klien.post("/api/vendor", json={"kode": "V009", "nama_baku": "PT X",
                                        "status": "aktif sekali"})
    assert r.status_code == 422
    assert "aktif" in r.json()["detail"]


def test_status_kosong_jadi_aktif(klien):
    r = klien.post("/api/vendor", json={"kode": "V010", "nama_baku": "PT X"})
    assert r.json()["vendor"]["status"] == "aktif"


def test_daftar_vendor_dan_pemakaian(klien_terpakai):
    d = klien_terpakai.get("/api/vendor").json()
    assert d["jumlah"] == 1 and d["jumlah_total"] == 1
    b = d["baris"][0]
    assert b["kode"] == "V001"
    assert b["jumlah_tiket"] == 1
    assert b["ton_kering"] == pytest.approx(12.0)
    assert b["baris_pembelian"] == 1
    assert b["intake_terakhir"] == "2026-01-05"
    assert b["terpakai"] is True


def test_daftar_cari_dan_saring_status(klien):
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Alpha",
                                    "kategori": "nasional"})
    klien.post("/api/vendor", json={"kode": "V002", "nama_baku": "CV Beta",
                                    "status": "nonaktif"})
    d = klien.get("/api/vendor", params={"q": "beta"}).json()
    assert [b["kode"] for b in d["baris"]] == ["V002"]
    d = klien.get("/api/vendor", params={"q": "nasional"}).json()   # cari kategori
    assert [b["kode"] for b in d["baris"]] == ["V001"]
    d = klien.get("/api/vendor", params={"status": "nonaktif"}).json()
    assert [b["kode"] for b in d["baris"]] == ["V002"]
    assert d["jumlah_total"] == 2          # total tidak ikut tersaring


def test_detail_vendor_tidak_ada_404(klien):
    assert klien.get("/api/vendor/V999").status_code == 404


# =================================================================== UPDATE
def test_ubah_vendor(klien):
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Alpha"})
    r = klien.put("/api/vendor/v001", json={
        "nama_baku": "PT Alpha Jaya", "varian_nama": "PT Alpha|CV Alpha",
        "kategori": "nasional", "status": "nonaktif", "pengisi": "budi"})
    assert r.status_code == 200, r.text
    v = r.json()["vendor"]
    assert v["nama_baku"] == "PT Alpha Jaya"
    assert v["varian_nama"] == "PT Alpha|CV Alpha"
    assert v["status"] == "nonaktif"
    d = klien.get("/api/vendor").json()
    assert d["baris"][0]["nama_baku"] == "PT Alpha Jaya"


def test_ubah_vendor_tidak_ada_404(klien):
    r = klien.put("/api/vendor/V404", json={"nama_baku": "PT X"})
    assert r.status_code == 404


def test_ubah_tidak_mengubah_riwayat_tiket(klien_terpakai):
    klien_terpakai.put("/api/vendor/V001", json={"nama_baku": "PT Alpha Baru"})
    d = klien_terpakai.get("/api/vendor").json()["baris"][0]
    assert d["nama_baku"] == "PT Alpha Baru"
    assert d["jumlah_tiket"] == 1          # tiket tetap menunjuk kode yang sama


def test_ubah_nama_kosong_memakai_nama_lama(klien):
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Alpha"})
    r = klien.put("/api/vendor/V001", json={"nama_baku": "   "})
    assert r.status_code == 200
    assert r.json()["vendor"]["nama_baku"] == "PT Alpha"


# ==================================================================== DELETE
def test_hapus_vendor_belum_terpakai(klien):
    klien.post("/api/vendor", json={"kode": "V005", "nama_baku": "PT Baru"})
    r = klien.delete("/api/vendor/V005", params={"pengisi": "budi"})
    assert r.status_code == 200
    assert "dihapus" in r.json()["pesan"]
    assert klien.get("/api/vendor").json()["jumlah"] == 0


def test_hapus_vendor_terpakai_ditolak(klien_terpakai):
    r = klien_terpakai.delete("/api/vendor/V001")
    assert r.status_code == 409
    pesan = r.json()["detail"]
    assert "nonaktif" in pesan          # menyarankan jalan yang aman
    assert "tiket" in pesan
    # vendor masih ada setelah penolakan
    assert klien_terpakai.get("/api/vendor/V001").status_code == 200


def test_hapus_paksa_vendor_terpakai(klien_terpakai):
    r = klien_terpakai.delete("/api/vendor/V001", params={"paksa": "true",
                                                          "pengisi": "budi"})
    assert r.status_code == 200
    assert r.json()["dipaksa"] is True
    assert r.json()["pemakaian"]["jumlah_tiket"] == 1
    # master-nya hilang ...
    assert klien_terpakai.get("/api/vendor/V001").status_code == 404
    # ... tetapi data tiket tetap ada: kartu tonase masih menghitung 12 ton V001
    ton = klien_terpakai.get("/api/kartu/tonase",
                             params={"dimensi": "vendor"}).json()
    assert any(b["kunci"] == "V001" and b["ton"] == pytest.approx(12.0)
               for b in ton["baris"])


def test_hapus_vendor_tidak_ada_404(klien):
    assert klien.delete("/api/vendor/V777").status_code == 404


# ================================================================== TEMPLATE
def test_template_bisa_dibuka_dan_berisi_kolom(klien):
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Alpha",
                                    "kategori": "nasional", "status": "aktif"})
    r = klien.get("/api/vendor/template")
    assert r.status_code == 200
    assert "Master_Vendor_template.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Master_Vendor" in wb.sheetnames
    assert "Petunjuk" in wb.sheetnames
    ws = wb["Master_Vendor"]
    assert [c.value for c in ws[1]] == ["Kode_Vendor", "Nama_Baku",
                                        "Varian_Nama", "Kategori", "Status",
                                        "Alamat", "Kontak_Person", "Klasifikasi"]
    assert ws["A2"].value == "V001"          # data yang ada ikut terisi


def test_template_kosong_beri_baris_contoh(klien):
    r = klien.get("/api/vendor/template")
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Master_Vendor"]
    assert ws.max_row == 2
    assert ws["A2"].value == "V001"


def test_template_bolak_balik_bisa_diunggah(klien):
    """Unduh template -> isi baris baru -> unggah kembali -> masuk."""
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Alpha"})
    isi = klien.get("/api/vendor/template").content
    wb = openpyxl.load_workbook(io.BytesIO(isi))
    ws = wb["Master_Vendor"]
    ws.append(["V042", "PT Hasil Bolak Balik", "PT HBB", "lokal", "aktif"])
    buf = io.BytesIO()
    wb.save(buf)
    r = _unggah(klien, buf.getvalue())
    assert r.status_code == 200, r.text
    assert r.json()["ditambah"] == 1
    assert r.json()["diperbarui"] == 1       # V001 dari template ikut diperbarui
    d = klien.get("/api/vendor").json()
    assert {b["kode"] for b in d["baris"]} == {"V001", "V042"}


# ============================================================ IMPOR MASSAL
def test_impor_menambah_dan_memperbarui(klien):
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Lama"})
    isi = _xlsx(baris=[["V001", "PT Baru", None, "nasional", "aktif"],
                       ["V002", "CV Beta", "CV B|CV Beta", "lokal", None]])
    r = _unggah(klien, isi)
    assert r.status_code == 200, r.text
    d = r.json()
    assert (d["ditambah"], d["diperbarui"], d["ditolak"]) == (1, 1, 0)
    daftar = {b["kode"]: b for b in klien.get("/api/vendor").json()["baris"]}
    assert daftar["V001"]["nama_baku"] == "PT Baru"
    assert daftar["V002"]["nama_baku"] == "CV Beta"
    assert daftar["V002"]["status"] == "aktif"      # kosong = aktif


def test_impor_baris_rusak_dilaporkan_tanpa_membatalkan_lain(klien):
    isi = _xlsx(baris=[
        ["V001", "PT Baik", None, None, "aktif"],     # lolos
        [None, "Tanpa Kode", None, None, None],       # kode kosong
        ["V002", None, None, None, None],             # nama kosong
        ["V#3", "Kode Aneh", None, None, None],       # kode tidak sah
        ["V004", "PT Status Salah", None, None, "entah"],   # status salah
        ["V005", "PT Baik 2", None, None, None],      # lolos
    ])
    r = _unggah(klien, isi)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["ditambah"] == 2
    assert d["ditolak"] == 4
    kode_ditolak = {p["kode"] for p in d["peringatan"]}
    assert kode_ditolak == {None, "V002", "V#3", "V004"}
    assert all("alasan" in p and "baris" in p for p in d["peringatan"])
    assert {b["kode"] for b in klien.get("/api/vendor").json()["baris"]} == \
        {"V001", "V005"}


def test_impor_kode_ganda_di_file_ditolak(klien):
    isi = _xlsx(baris=[["V001", "PT Satu", None, None, "aktif"],
                       ["v001", "PT Dua", None, None, "aktif"]])
    d = _unggah(klien, isi).json()
    assert d["ditambah"] == 1
    assert d["ditolak"] == 1
    assert "ganda" in d["peringatan"][0]["alasan"].lower()
    d2 = klien.get("/api/vendor").json()["baris"][0]
    assert d2["nama_baku"] == "PT Satu"      # baris pertama yang menang


def test_impor_bukan_xlsx_ditolak(klien):
    r = klien.post("/api/vendor/impor",
                   files={"file": ("vendor.csv", b"a,b", "text/csv")})
    assert r.status_code == 422
    assert ".xlsx" in r.json()["detail"]


def test_impor_tanpa_sheet_master_vendor_ditolak(klien):
    isi = _xlsx(sheet="Sheet1", baris=[["V001", "PT X", None, None, None]])
    r = _unggah(klien, isi)
    assert r.status_code == 422
    assert "Master_Vendor" in r.json()["detail"]


def test_impor_file_rusak_ditolak(klien):
    r = _unggah(klien, b"bukan excel sama sekali", "rusak.xlsx")
    assert r.status_code == 422


def test_impor_tidak_mengubah_yang_tidak_ada_di_file(klien):
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Tetap",
                                    "kategori": "nasional"})
    _unggah(klien, _xlsx(baris=[["V002", "CV Baru", None, None, None]]))
    daftar = {b["kode"]: b for b in klien.get("/api/vendor").json()["baris"]}
    assert daftar["V001"]["nama_baku"] == "PT Tetap"
    assert daftar["V001"]["kategori"] == "nasional"


# ================================================ PROFIL VENDOR (v0.10.0)
# Permintaan pemilik: "tambahkan informasi seperti Address, Contact person,
# klasifikasi sebagai Small holder, small estate, atau Rubber trader."


def test_tambah_vendor_dengan_profil(klien):
    r = klien.post("/api/vendor", json={
        "kode": "v010", "nama_baku": "CV Profil", "pengisi": "budi",
        "alamat": "  Jl. Gatot Subroto No. 12, Medan  ",
        "kontak_person": "Budi — 0812-1111-2222",
        "klasifikasi": "Small holder",
    })
    assert r.status_code == 200, r.text
    v = r.json()["vendor"]
    assert v["alamat"] == "Jl. Gatot Subroto No. 12, Medan"     # tepi dibersihkan
    assert v["kontak_person"] == "Budi — 0812-1111-2222"
    assert v["klasifikasi"] == "small_holder"                   # disimpan sebagai kode
    assert v["klasifikasi_label"] == "Small holder"             # tampil sebagai label


@pytest.mark.parametrize("ejaan,kode", [
    ("Small holder", "small_holder"),
    ("smallholder", "small_holder"),
    ("SMALL HOLDER", "small_holder"),
    ("Kebun Rakyat", "small_holder"),
    ("Small estate", "small_estate"),
    ("small-estate", "small_estate"),
    ("Rubber trader", "rubber_trader"),
    ("trader", "rubber_trader"),
    ("Pedagang", "rubber_trader"),
])
def test_ejaan_klasifikasi_diterima(klien, ejaan, kode):
    r = klien.post("/api/vendor", json={"kode": "V1", "nama_baku": "PT Satu",
                                        "klasifikasi": ejaan})
    assert r.status_code == 200, r.text
    assert r.json()["vendor"]["klasifikasi"] == kode


def test_klasifikasi_tidak_dikenal_ditolak(klien):
    r = klien.post("/api/vendor", json={"kode": "V1", "nama_baku": "PT Satu",
                                        "klasifikasi": "koperasi besar"})
    assert r.status_code == 422
    detail = r.json()["detail"]
    assert "Small holder" in detail and "Rubber trader" in detail
    assert klien.get("/api/vendor").json()["baris"] == []


def test_klasifikasi_kosong_boleh(klien):
    r = klien.post("/api/vendor", json={"kode": "V1", "nama_baku": "PT Satu",
                                        "klasifikasi": ""})
    assert r.status_code == 200
    assert r.json()["vendor"]["klasifikasi"] is None
    assert r.json()["vendor"]["klasifikasi_label"] is None


def test_ubah_profil_dan_audit_menyebut_nilai_lama_baru(klien):
    klien.post("/api/vendor", json={"kode": "V1", "nama_baku": "PT Satu",
                                    "alamat": "Jl. Lama"})
    r = klien.put("/api/vendor/V1", json={
        "nama_baku": "PT Satu", "alamat": "Jl. Baru No. 5",
        "kontak_person": "Siti", "klasifikasi": "rubber_trader"})
    assert r.status_code == 200, r.text
    v = r.json()["vendor"]
    assert (v["alamat"], v["kontak_person"], v["klasifikasi"]) == \
        ("Jl. Baru No. 5", "Siti", "rubber_trader")
    # mengosongkan = menghapus isian (bukan menyimpan string kosong)
    r2 = klien.put("/api/vendor/V1", json={"nama_baku": "PT Satu",
                                           "alamat": "", "kontak_person": "",
                                           "klasifikasi": ""})
    v2 = r2.json()["vendor"]
    assert (v2["alamat"], v2["kontak_person"], v2["klasifikasi"]) == (None, None, None)


def test_daftar_saring_klasifikasi_dan_ringkasan(klien):
    for kode, kelas in [("V1", "small_holder"), ("V2", "small_holder"),
                        ("V3", "small_estate"), ("V4", "rubber_trader"),
                        ("V5", "")]:
        klien.post("/api/vendor", json={"kode": kode, "nama_baku": f"PT {kode}",
                                        "klasifikasi": kelas})
    semua = klien.get("/api/vendor").json()
    assert semua["ringkasan_klasifikasi"] == {
        "small_holder": 2, "small_estate": 1, "rubber_trader": 1, "belum_diisi": 1}
    assert [k["kode"] for k in semua["klasifikasi_baku"]] == \
        ["small_holder", "small_estate", "rubber_trader"]

    saring = klien.get("/api/vendor?klasifikasi=small_holder").json()
    assert {b["kode"] for b in saring["baris"]} == {"V1", "V2"}
    # saringan menerima label maupun kode, tidak peka huruf besar/kecil
    assert {b["kode"] for b in
            klien.get("/api/vendor?klasifikasi=Rubber%20trader").json()["baris"]} == {"V4"}
    assert [b["kode"] for b in
            klien.get("/api/vendor?klasifikasi=-").json()["baris"]] == ["V5"]
    assert klien.get("/api/vendor?klasifikasi=entah").status_code == 422


def test_cari_menemukan_alamat_kontak_dan_klasifikasi(klien):
    klien.post("/api/vendor", json={"kode": "V1", "nama_baku": "PT Alpha",
                                    "alamat": "Jl. Medan Baru", "kontak_person": "Siti",
                                    "klasifikasi": "rubber_trader"})
    assert [b["kode"] for b in klien.get("/api/vendor?q=medan").json()["baris"]] == ["V1"]
    assert [b["kode"] for b in klien.get("/api/vendor?q=siti").json()["baris"]] == ["V1"]
    assert [b["kode"] for b in klien.get("/api/vendor?q=trader").json()["baris"]] == ["V1"]


def test_template_memuat_kolom_profil_dan_petunjuk(klien):
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Alpha",
                                    "alamat": "Jl. Contoh 9",
                                    "kontak_person": "Budi", "klasifikasi": "small_estate"})
    wb = openpyxl.load_workbook(io.BytesIO(klien.get("/api/vendor/template").content))
    ws = wb["Master_Vendor"]
    assert [c.value for c in ws[1]][-3:] == ["Alamat", "Kontak_Person", "Klasifikasi"]
    assert ws["F2"].value == "Jl. Contoh 9"
    assert ws["G2"].value == "Budi"
    assert ws["H2"].value == "small_estate"
    petunjuk = "\n".join(str(c.value or "") for (c,) in wb["Petunjuk"].iter_rows())
    assert "Small holder" in petunjuk and "Rubber trader" in petunjuk


def _xlsx_profil(baris: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master_Vendor"
    ws.append(["Kode_Vendor", "Nama_Baku", "Alamat", "Kontak_Person", "Klasifikasi"])
    for b in baris:
        ws.append(b)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_impor_profil_alias_diterima(klien):
    isi = _xlsx_profil([
        ["V001", "PT Alpha", "Jl. Satu 1", "Budi — 0812", "Smallholder"],
        ["V002", "CV Beta", None, None, "Rubber Trader"],
        ["V003", "PT Gamma", "Jl. Tiga 3", "Siti", "small estate"],
    ])
    d = _unggah(klien, isi).json()
    assert (d["ditambah"], d["diperbarui"], d["ditolak"]) == (3, 0, 0)
    daftar = {b["kode"]: b for b in klien.get("/api/vendor").json()["baris"]}
    assert daftar["V001"]["klasifikasi"] == "small_holder"
    assert daftar["V001"]["alamat"] == "Jl. Satu 1"
    assert daftar["V001"]["kontak_person"] == "Budi — 0812"
    assert daftar["V002"]["klasifikasi"] == "rubber_trader"
    assert daftar["V003"]["klasifikasi"] == "small_estate"


def test_impor_klasifikasi_salah_ditolak_per_baris(klien):
    isi = _xlsx_profil([
        ["V001", "PT Baik", None, None, "small_holder"],
        ["V002", "PT Salah", None, None, "koperasi"],
        ["V003", "PT Baik 2", None, None, None],
    ])
    d = _unggah(klien, isi).json()
    assert d["ditambah"] == 2
    assert d["ditolak"] == 1
    salah = [p for p in d["peringatan"] if p["kode"] == "V002"][0]
    assert "Small holder" in salah["alasan"]
    assert {b["kode"] for b in klien.get("/api/vendor").json()["baris"]} == \
        {"V001", "V003"}


def test_impor_tidak_menghapus_profil_yang_tidak_ada_di_file(klien):
    """File berisi kode+nama saja tidak boleh menghapus alamat/kontak."""
    klien.post("/api/vendor", json={"kode": "V001", "nama_baku": "PT Alpha",
                                    "alamat": "Jl. Penting", "kontak_person": "Budi",
                                    "klasifikasi": "small_holder"})
    isi = _xlsx(baris=[["V001", "PT Alpha Baru", None, None, "aktif"]])
    assert _unggah(klien, isi).json()["diperbarui"] == 1
    v = klien.get("/api/vendor").json()["baris"][0]
    assert v["nama_baku"] == "PT Alpha Baru"
    assert v["alamat"] == "Jl. Penting"
    assert v["kontak_person"] == "Budi"
    assert v["klasifikasi"] == "small_holder"


def test_impor_template_utama_juga_mengisi_profil():
    """Sheet Master_Vendor di template utama (Unggah Data) ikut mengisi profil."""
    import tempfile

    from app.importer.template import impor_template
    from sqlalchemy import create_engine, select
    from sqlalchemy.orm import sessionmaker

    from app.models import Base, Vendor

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master_Vendor"
    ws.append(["Kode_Vendor", "Nama_Baku", "Varian_Nama", "Kategori", "Status",
               "Alamat", "Kontak_Person", "Klasifikasi"])
    ws.append(["V100", "PT Profil", None, "lokal", "aktif",
               "Jl. Uji 10", "Rina — 0813", "Small Estate"])
    ws.append(["V101", "PT Salah Kelas", None, None, None, None, None, "entah"])
    with tempfile.TemporaryDirectory() as tmp:
        path = f"{tmp}/profil.xlsx"
        wb.save(path)
        eng = create_engine(f"sqlite:///{tmp}/p.db", future=True)
        Base.metadata.create_all(eng)
        with sessionmaker(bind=eng, future=True)() as s:
            hasil = impor_template(path, s, pengunggah="pytest")
            assert hasil.jumlah["vendor"] == 1          # V101 ditolak
            assert hasil.ditolak == {"Master_Vendor": 1}
            v = s.scalars(select(Vendor).where(Vendor.kode == "V100")).one()
            assert v.klasifikasi == "small_estate"
            assert v.kontak_person == "Rina — 0813"
            assert v.alamat == "Jl. Uji 10"
