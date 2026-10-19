"""Tes Fase 7: master grade (menu isi grade) & harga per grade (form + bulk)."""
import datetime as dt
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool

from app.main import buat_app


@pytest.fixture()
def klien(tmp_path, monkeypatch):
    monkeypatch.setenv("UNGGAH_DIR", str(tmp_path / "unggahan"))
    eng = create_engine("sqlite://", future=True,
                        connect_args={"check_same_thread": False},
                        poolclass=StaticPool)
    k = TestClient(buat_app(eng))
    # data dasar: kurs & sicom supaya konversi US¢ bisa dihitung
    k.post("/api/kurs", json={"berlaku_dari": "2026-10-01", "idr_per_usd": 16000})
    k.post("/api/sicom", json={"tanggal": "2026-10-16", "usc_per_kg": 250.0})
    return k


def test_master_grade_tambah_ubah_dan_audit(klien):
    r = klien.post("/api/grade", json={
        "item_code": "SLAB2", "kelompok": "Slab 2", "eudr": False,
        "keterangan": "grade utama", "pengisi": "admin-uji"})
    assert r.status_code == 200 and "ditambahkan" in r.json()["pesan"]
    # kode TTSR otomatis EUDR walau tidak dicentang
    klien.post("/api/grade", json={"item_code": "SLAB2 TTSR", "pengisi": "admin-uji"})
    daftar = klien.get("/api/grade").json()
    m = {g["item_code"]: g for g in daftar}
    assert m["SLAB2"]["kelompok"] == "Slab 2" and not m["SLAB2"]["eudr"]
    assert m["SLAB2 TTSR"]["eudr"]
    # ubah
    r = klien.post("/api/grade", json={
        "item_code": "SLAB2", "kelompok": "Slab 2 rev", "pengisi": "admin-uji"})
    assert "diperbarui" in r.json()["pesan"]
    audit = klien.get("/api/audit", params={"objek": "grade"}).json()
    assert audit and audit[0]["siapa"] == "admin-uji"
    assert audit[0]["detail"]["lama"]["kelompok"] == "Slab 2"


def test_harga_grade_form_dan_terkini(klien):
    klien.post("/api/grade", json={"item_code": "SLAB2"})
    klien.post("/api/grade", json={"item_code": "SLAB3"})
    r = klien.post("/api/harga-grade", json={
        "tanggal": "2026-10-15",
        "baris": [{"item_code": "SLAB2", "harga_rp_per_kg": 28000},
                  {"item_code": "SLAB3", "harga_rp_per_kg": 27000}],
        "pengisi": "staf-uji"})
    assert r.status_code == 200 and "2 harga grade" in r.json()["pesan"]

    t = klien.get("/api/harga-grade/terkini", params={"tanggal": "2026-10-16"}).json()
    m = {b["item_code"]: b for b in t["baris"]}
    # harga berlaku: baris terbaru <= tanggal (15 Okt berlaku utk 16 Okt)
    assert m["SLAB2"]["harga_rp_per_kg"] == 28000
    assert m["SLAB2"]["harga_usc_per_kg"] == round(28000 / 16000 * 100, 4)  # 175
    assert m["SLAB2"]["margin_indikatif_usc"] == round(250 - 175, 4)
    assert "INDIKATIF" in t["catatan"]

    # timpa tanggal sama -> tidak menggandakan
    klien.post("/api/harga-grade", json={
        "tanggal": "2026-10-15",
        "baris": [{"item_code": "SLAB2", "harga_rp_per_kg": 28500}]})
    riwayat = klien.get("/api/harga-grade").json()
    slab2 = [b for b in riwayat if b["item_code"] == "SLAB2"]
    assert len(slab2) == 1 and slab2[0]["harga_rp_per_kg"] == 28500


def test_harga_grade_tolak_grade_tak_dikenal(klien):
    r = klien.post("/api/harga-grade", json={
        "tanggal": "2026-10-15",
        "baris": [{"item_code": "GRADE-X", "harga_rp_per_kg": 1000}]})
    assert r.status_code == 422
    assert "tidak ada di master" in r.json()["detail"]


def test_harga_grade_tolak_harga_nol(klien):
    klien.post("/api/grade", json={"item_code": "SLAB2"})
    r = klien.post("/api/harga-grade", json={
        "tanggal": "2026-10-15",
        "baris": [{"item_code": "SLAB2", "harga_rp_per_kg": 0}]})
    assert r.status_code == 422


def test_template_unduh_lalu_bulk_import(klien):
    """Alur bulk: unduh template -> isi harga -> unggah di tab Unggah Data."""
    klien.post("/api/grade", json={"item_code": "SLAB2"})
    klien.post("/api/grade", json={"item_code": "LATEX OP"})
    r = klien.get("/api/harga-grade/template")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Harga_Grade"]
    ws = wb["Harga_Grade"]
    assert ws.cell(1, 1).value == "Tanggal"
    # isi harga seperti yang akan dilakukan pengguna
    for baris in ws.iter_rows(min_row=2):
        baris[2].value = 26000
    buf = io.BytesIO(); wb.save(buf)

    r = klien.post("/api/unggah",
                   files={"file": ("harga.xlsx", buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument"
                          ".spreadsheetml.sheet")},
                   data={"pengisi": "staf-uji"})
    assert r.status_code == 200
    assert r.json()["jumlah"]["harga_grade"] == 2
    t = klien.get("/api/harga-grade/terkini").json()
    m = {b["item_code"]: b for b in t["baris"]}
    assert m["SLAB2"]["harga_rp_per_kg"] == 26000


def test_bulk_import_idempoten_dan_validasi(klien):
    klien.post("/api/grade", json={"item_code": "SLAB2"})
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Harga_Grade"
    ws.append(["Tanggal", "Item_Code", "Harga_Rp_per_Kg", "Catatan"])
    ws.append([dt.date(2026, 10, 15), "SLAB2", 28000, None])
    ws.append([dt.date(2026, 10, 15), "GRADE-ANEH", 1000, None])   # warning
    ws.append(["bukan-tanggal", "SLAB2", 28000, None])             # ditolak
    buf = io.BytesIO(); wb.save(buf)
    kirim = lambda: klien.post("/api/unggah",
        files={"file": ("h.xlsx", buf.getvalue(), "application/x")},
        data={"pengisi": "x"})
    j = kirim().json()
    assert j["jumlah"]["harga_grade"] == 2  # baris sah (termasuk grade asing)
    assert j["ditolak"].get("Harga_Grade") == 1
    teks = "\n".join(j["peringatan"])
    assert "GRADE-ANEH" in teks and "bukan-tanggal" in teks
    j2 = kirim().json()
    assert j2["duplikat_file"]
    assert len(klien.get("/api/harga-grade").json()) == 2  # tidak berganda
