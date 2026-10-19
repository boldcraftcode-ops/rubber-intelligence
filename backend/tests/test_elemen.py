"""Tes kartu 'RM Stock-Intake VS Pricing' (disamakan dengan Excel lampiran)."""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import buat_app


@pytest.fixture(scope="module")
def klien(engine):
    return TestClient(buat_app(engine))


def _baris(data, label):
    return next(b for b in data["baris"] if b["label"] == label)


def test_cocok_dengan_excel_lampiran(klien):
    d = klien.get("/api/ringkasan-elemen").json()

    jan23 = _baris(d, "Jan-23")
    assert jan23["rm_stock_ton"] == 4313.93
    assert jan23["intake_ton"] == 7894.32
    assert jan23["safety_ton"] == 3850
    assert round(jan23["sicom"], 4) == 138.6176
    assert round(jan23["op_price"], 4) == 122.2813

    mei24 = _baris(d, "Mei-24")
    assert mei24["rm_stock_ton"] == 2915.551
    assert mei24["safety_ton"] == 2795.1  # safety berubah per Mei-24

    agu26 = _baris(d, "Agu-26")
    assert agu26["rm_stock_ton"] == 551.947
    # Intake = tonase tiket (sumber resmi) dari referensi rekonsiliasi
    assert agu26["intake_ton"] == 4759.827
    assert "tiket" in agu26["sumber_intake"]
    jul26 = _baris(d, "Jul-26")
    assert jul26["intake_ton"] == 5269.319


def test_jun25_memakai_data_harian_bukan_excel(klien):
    """Excel lampiran menulis 2.795,1 (sama persis dengan safety) — diduga
    salah isi; sistem memakai stok harian 30/06/2025 dan memberi catatan."""
    d = klien.get("/api/ringkasan-elemen").json()
    jun25 = _baris(d, "Jun-25")
    assert jun25["rm_stock_ton"] == 2038.681
    assert any("Jun-25" in c for c in d["catatan"])


def test_tanpa_baris_proyeksi(klien):
    """Baris proyeksi Sep-Des 26 di Excel tidak direplikasi; bulan berjalan
    dengan data stok belum penuh ditandai 'sebagian'."""
    d = klien.get("/api/ringkasan-elemen").json()
    sep26 = _baris(d, "Sep-26")
    assert sep26["stok_sebagian"]  # stok terakhir 08/09, bukan akhir bulan
    assert all(not b["label"].endswith("Nov-26") for b in d["baris"]
               if b["rm_stock_ton"] is None and b["intake_ton"] is None)


def test_ekspor_xlsx(klien):
    r = klien.get("/api/ringkasan-elemen", params={"format": "xlsx"})
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1][:6]] == \
        ["Date", "RM Stock", "Intake", "Safety Level Stock", "SICOM", "OP Price"]
    assert ws.max_row >= 44
