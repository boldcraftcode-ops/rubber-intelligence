"""API Master Collector: CRUD, template Excel, dan impor massal.

Collector adalah perorangan yang terafiliasi ke satu vendor. Wilayah
operasional dicatat terpisah sebagai Provinsi, Kabupaten/Kota, dan Desa.
"""
from __future__ import annotations

import io
import re

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .core.alarm import catat_audit
from .models import Collector, Vendor

SHEET = "Master_Collector"
KOLOM = ["Kode_Kolektor", "Nama_Kolektor", "NIK", "No_HP", "Afiliasi_Vendor",
         "Provinsi", "Kabupaten_Kota", "Desa"]
POLA_KODE = re.compile(r"^[A-Z0-9][A-Z0-9._/-]{0,19}$")
POLA_NIK = re.compile(r"^\d{16}$")
BATAS_IMPOR = 5000
BATAS_LAPORAN = 200


class CollectorMasuk(BaseModel):
    kode: str = Field(min_length=1, max_length=20)
    nama: str = Field(min_length=1, max_length=200)
    nik: str | None = None
    no_hp: str | None = Field(default=None, max_length=30)
    kode_vendor: str = Field(min_length=1, max_length=20)
    provinsi: str | None = Field(default=None, max_length=100)
    kabupaten_kota: str | None = Field(default=None, max_length=120)
    desa: str | None = Field(default=None, max_length=120)
    pengisi: str = "?"


class CollectorUbah(BaseModel):
    nama: str = Field(min_length=1, max_length=200)
    nik: str | None = None
    no_hp: str | None = Field(default=None, max_length=30)
    kode_vendor: str = Field(min_length=1, max_length=20)
    provinsi: str | None = Field(default=None, max_length=100)
    kabupaten_kota: str | None = Field(default=None, max_length=120)
    desa: str | None = Field(default=None, max_length=120)
    pengisi: str = "?"


def _kode(teks: str | None) -> str:
    return (teks or "").strip().upper()


def _bersih(teks: str | None) -> str | None:
    nilai = (teks or "").strip()
    return nilai or None


def _kode_valid(teks: str | None) -> str:
    hasil = _kode(teks)
    if not hasil:
        raise HTTPException(422, "Kode kolektor wajib diisi.")
    if not POLA_KODE.fullmatch(hasil):
        raise HTTPException(422, "Kode kolektor hanya boleh berisi huruf, angka, "
                                 "titik, garis bawah, garis miring, atau tanda hubung; "
                                 "maksimal 20 karakter.")
    return hasil


def _nik_valid(teks: str | None) -> str | None:
    hasil = _bersih(teks)
    if hasil and not POLA_NIK.fullmatch(hasil):
        raise HTTPException(422, "NIK harus tepat 16 angka. Simpan sebagai teks "
                                 "di Excel agar angka tidak berubah.")
    return hasil


def _periksa_vendor(session: Session, kode_vendor: str | None) -> tuple[str, Vendor]:
    kode = _kode(kode_vendor)
    if not kode:
        raise HTTPException(422, "Afiliasi Vendor wajib dipilih.")
    vendor = session.get(Vendor, kode)
    if vendor is None:
        raise HTTPException(422, f"Vendor '{kode}' belum ada di Master Vendor. "
                                 "Tambahkan vendor terlebih dahulu.")
    return kode, vendor


def _periksa_nik(session: Session, nik: str | None, kode: str) -> str | None:
    nik = _nik_valid(nik)
    if nik and session.scalar(select(Collector.kode).where(
            Collector.nik == nik, Collector.kode != kode)):
        raise HTTPException(409, "NIK tersebut sudah terdaftar pada kolektor lain.")
    return nik


def _baris(session: Session, c: Collector) -> dict:
    vendor = session.get(Vendor, c.kode_vendor)
    return {
        "kode": c.kode, "nama": c.nama, "nik": c.nik, "no_hp": c.no_hp,
        "kode_vendor": c.kode_vendor,
        "nama_vendor": vendor.nama_baku if vendor else "(vendor tidak ditemukan)",
        "provinsi": c.provinsi, "kabupaten_kota": c.kabupaten_kota,
        "desa": c.desa,
    }


def pasang_router_collector(app, ambil_session):
    router = APIRouter(prefix="/api")
    dep = Depends(ambil_session)

    @router.get("/collector")
    def daftar(q: str | None = None, kode_vendor: str | None = None,
               provinsi: str | None = None, kabupaten_kota: str | None = None,
               desa: str | None = None, session: Session = dep):
        stmt = select(Collector).order_by(Collector.kode)
        if kode_vendor:
            stmt = stmt.where(Collector.kode_vendor == _kode(kode_vendor))
        for kolom, nilai in ((Collector.provinsi, provinsi),
                             (Collector.kabupaten_kota, kabupaten_kota),
                             (Collector.desa, desa)):
            if nilai and nilai.strip():
                stmt = stmt.where(func.lower(kolom) == nilai.strip().lower())
        baris = [_baris(session, c) for c in session.scalars(stmt)]
        cari = (q or "").strip().lower()
        if cari:
            baris = [b for b in baris if any(
                cari in (str(b[k] or "").lower()) for k in (
                    "kode", "nama", "nik", "no_hp", "kode_vendor",
                    "nama_vendor", "provinsi", "kabupaten_kota", "desa"))]
        return {"baris": baris, "jumlah": len(baris),
                "jumlah_total": session.scalar(select(func.count()).select_from(Collector)) or 0}

    @router.post("/collector")
    def tambah(data: CollectorMasuk, session: Session = dep):
        kode = _kode_valid(data.kode)
        if session.get(Collector, kode):
            raise HTTPException(409, f"Kode kolektor '{kode}' sudah terdaftar.")
        kode_vendor, _ = _periksa_vendor(session, data.kode_vendor)
        nik = _periksa_nik(session, data.nik, kode)
        nama = _bersih(data.nama)
        if not nama:
            raise HTTPException(422, "Nama kolektor wajib diisi.")
        c = Collector(kode=kode, nama=nama, nik=nik, no_hp=_bersih(data.no_hp),
                       kode_vendor=kode_vendor, provinsi=_bersih(data.provinsi),
                       kabupaten_kota=_bersih(data.kabupaten_kota), desa=_bersih(data.desa))
        session.add(c)
        catat_audit(session, data.pengisi, "tambah", "collector", {
            "kode": kode, "nama": nama, "kode_vendor": kode_vendor,
            "wilayah": [c.provinsi, c.kabupaten_kota, c.desa]})
        session.commit()
        return {"status": "tersimpan", "pesan": f"Kolektor {kode} — {nama} ditambahkan.",
                "collector": _baris(session, c)}

    @router.put("/collector/{kode}")
    def ubah(kode: str, data: CollectorUbah, session: Session = dep):
        kode = _kode_valid(kode)
        c = session.get(Collector, kode)
        if c is None:
            raise HTTPException(404, f"Kolektor '{kode}' tidak ditemukan.")
        kode_vendor, _ = _periksa_vendor(session, data.kode_vendor)
        nik = _periksa_nik(session, data.nik, kode)
        lama = {"nama": c.nama, "nik": c.nik, "no_hp": c.no_hp,
                "kode_vendor": c.kode_vendor, "provinsi": c.provinsi,
                "kabupaten_kota": c.kabupaten_kota, "desa": c.desa}
        c.nama = _bersih(data.nama) or c.nama
        c.nik = nik
        c.no_hp = _bersih(data.no_hp)
        c.kode_vendor = kode_vendor
        c.provinsi = _bersih(data.provinsi)
        c.kabupaten_kota = _bersih(data.kabupaten_kota)
        c.desa = _bersih(data.desa)
        catat_audit(session, data.pengisi, "ubah", "collector", {
            "kode": kode, "lama": lama,
            "baru": {"nama": c.nama, "nik": c.nik, "no_hp": c.no_hp,
                    "kode_vendor": c.kode_vendor, "provinsi": c.provinsi,
                    "kabupaten_kota": c.kabupaten_kota, "desa": c.desa}})
        session.commit()
        return {"status": "tersimpan", "pesan": f"Kolektor {kode} diperbarui.",
                "collector": _baris(session, c)}

    @router.delete("/collector/{kode}")
    def hapus(kode: str, pengisi: str = "?", session: Session = dep):
        kode = _kode_valid(kode)
        c = session.get(Collector, kode)
        if c is None:
            raise HTTPException(404, f"Kolektor '{kode}' tidak ditemukan.")
        detail = {"kode": c.kode, "nama": c.nama, "kode_vendor": c.kode_vendor}
        session.delete(c)
        catat_audit(session, pengisi, "hapus", "collector", detail)
        session.commit()
        return {"status": "terhapus", "pesan": f"Kolektor {kode} dihapus."}

    @router.get("/collector/template")
    def template(session: Session = dep):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(KOLOM)
        kolektor = list(session.scalars(select(Collector).order_by(Collector.kode)))
        for c in kolektor:
            ws.append([c.kode, c.nama, c.nik, c.no_hp, c.kode_vendor,
                       c.provinsi, c.kabupaten_kota, c.desa])
        if not kolektor:
            vendor = session.scalar(select(Vendor).order_by(Vendor.kode))
            ws.append(["K001", "Budi Contoh", "3201010101900001", "081200000001",
                       vendor.kode if vendor else "V001", "Sumatera Utara",
                       "Kota Medan", "Medan Sunggal"])
        for kolom, lebar in zip("ABCDEFGH", (18, 32, 22, 20, 22, 24, 28, 28)):
            ws.column_dimensions[kolom].width = lebar
        ws.freeze_panes = "A2"
        panduan = wb.create_sheet("Petunjuk")
        for row in [
            ["PETUNJUK MASTER COLLECTOR"], [""],
            ["1. Isi data pada sheet Master_Collector; jangan ubah nama sheet."],
            ["2. Kode_Kolektor, Nama_Kolektor, dan Afiliasi_Vendor wajib."],
            ["3. Afiliasi_Vendor harus memakai Kode_Vendor yang ada di Master Vendor."],
            ["4. NIK harus 16 angka. Atur format sel NIK sebagai Teks sebelum mengisi."],
            ["5. Wilayah diisi bertingkat: Provinsi, Kabupaten/Kota, lalu Desa."],
            ["6. Kode kolektor yang sudah ada akan diperbarui, bukan digandakan."],
            ["7. Sel opsional kosong saat impor tidak menghapus data yang sudah ada."],
            ["8. Baris yang salah dilaporkan; baris lain yang benar tetap diproses."],
        ]:
            panduan.append(row)
        panduan.column_dimensions["A"].width = 115
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return StreamingResponse(out,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="Template_Master_Collector.xlsx"'})

    @router.post("/collector/impor")
    async def impor(file: UploadFile = File(...), pengisi: str = Form("?"),
                    session: Session = dep):
        if not (file.filename or "").lower().endswith(".xlsx"):
            raise HTTPException(422, "Pilih file Excel .xlsx.")
        isi = await file.read()
        try:
            wb = openpyxl.load_workbook(io.BytesIO(isi), read_only=True, data_only=True)
        except Exception as e:
            raise HTTPException(422, "File Excel tidak dapat dibaca.") from e
        ws = wb[SHEET] if SHEET in wb.sheetnames else None
        if ws is None:
            wb.close()
            raise HTTPException(422, f"Sheet '{SHEET}' tidak ditemukan.")
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            raise HTTPException(422, "Sheet Master_Collector kosong.")
        header_ix = next((i for i, row in enumerate(rows[:10])
                          if any(str(x or "").strip().lower().replace(" ", "_")
                                 in ("kode_kolektor", "kode_collector") for x in row)), None)
        if header_ix is None:
            raise HTTPException(422, "Header Kode_Kolektor tidak ditemukan.")
        mapping: dict[str, int] = {}
        alias = {
            "kode": {"kode_kolektor", "kode_collector", "kode"},
            "nama": {"nama_kolektor", "nama_collector", "nama"},
            "nik": {"nik", "nomor_induk_kependudukan"},
            "no_hp": {"no_hp", "no._hp", "nomor_hp", "telepon", "phone"},
            "vendor": {"afiliasi_vendor", "kode_vendor", "vendor"},
            "provinsi": {"provinsi", "province"},
            "kabupaten_kota": {"kabupaten_kota", "kabupaten/kota", "kabupaten", "kota"},
            "desa": {"desa", "village"},
        }
        for i, cell in enumerate(rows[header_ix]):
            k = str(cell or "").strip().lower().replace(" ", "_")
            for field, names in alias.items():
                if k in names:
                    mapping[field] = i
        wajib = {"kode", "nama", "vendor"}
        if not wajib <= mapping.keys():
            raise HTTPException(422, "Kolom Kode_Kolektor, Nama_Kolektor, dan "
                                     "Afiliasi_Vendor wajib ada.")
        data_rows = [(no, row) for no, row in enumerate(rows[header_ix + 1:], header_ix + 2)
                     if any(x not in (None, "") for x in row)]
        if len(data_rows) > BATAS_IMPOR:
            raise HTTPException(422, f"File melebihi batas {BATAS_IMPOR} baris sekali impor.")

        def val(row, field):
            i = mapping.get(field)
            if i is None or i >= len(row) or row[i] is None:
                return None
            return str(row[i]).strip() or None

        ditambah = diperbarui = 0
        tolak: list[dict] = []
        dilihat: dict[str, int] = {}
        for no, row in data_rows:
            kode = _kode(val(row, "kode"))
            nama = _bersih(val(row, "nama"))
            vendor_code = _kode(val(row, "vendor"))
            reason = None
            if not kode:
                reason = "Kode_Kolektor kosong."
            elif not POLA_KODE.fullmatch(kode):
                reason = "Kode kolektor tidak valid (maksimal 20 karakter)."
            elif not nama:
                reason = "Nama_Kolektor kosong."
            elif len(nama) > 200:
                reason = "Nama_Kolektor melebihi 200 karakter."
            elif not vendor_code:
                reason = "Afiliasi_Vendor wajib diisi."
            elif kode in dilihat:
                reason = f"Kode ganda di file (sudah dipakai baris {dilihat[kode]})."
            elif session.get(Vendor, vendor_code) is None:
                reason = f"Vendor '{vendor_code}' belum ada di Master Vendor."
            nik = val(row, "nik")
            if not reason:
                try:
                    nik = _nik_valid(nik)
                except HTTPException as e:
                    reason = str(e.detail)
            if not reason:
                for field, limit, label in (("no_hp", 30, "No_HP"),
                                             ("provinsi", 100, "Provinsi"),
                                             ("kabupaten_kota", 120, "Kabupaten_Kota"),
                                             ("desa", 120, "Desa")):
                    nilai = _bersih(val(row, field))
                    if nilai and len(nilai) > limit:
                        reason = f"{label} melebihi {limit} karakter."
                        break
            if not reason and nik and session.scalar(select(Collector.kode).where(
                    Collector.nik == nik, Collector.kode != kode)): 
                reason = "NIK tersebut sudah terdaftar pada kolektor lain."
            if reason:
                tolak.append({"baris": no, "kode": kode or None, "alasan": reason})
                continue
            dilihat[kode] = no
            c = session.get(Collector, kode)
            if c is None:
                c = Collector(kode=kode, nama=nama, kode_vendor=vendor_code)
                session.add(c)
                ditambah += 1
            else:
                c.nama = nama
                c.kode_vendor = vendor_code
                diperbarui += 1
            # Kolom opsional kosong = jangan menghapus nilai lama.
            for field, attr in (("nik", "nik"), ("no_hp", "no_hp"),
                                ("provinsi", "provinsi"),
                                ("kabupaten_kota", "kabupaten_kota"), ("desa", "desa")):
                nilai = _bersih(val(row, field))
                if nilai is not None:
                    setattr(c, attr, nilai)
            session.flush()
        catat_audit(session, pengisi, "impor", "collector", {
            "nama_file": file.filename, "ditambah": ditambah,
            "diperbarui": diperbarui, "ditolak": len(tolak),
            "kode": sorted(dilihat)})
        session.commit()
        return {"status": "selesai", "nama_file": file.filename,
                "ditambah": ditambah, "diperbarui": diperbarui,
                "ditolak": len(tolak), "jumlah_baris": len(data_rows),
                "peringatan": tolak[:BATAS_LAPORAN],
                "pesan": f"{ditambah} kolektor baru, {diperbarui} diperbarui"
                         + (f", {len(tolak)} baris ditolak" if tolak else "") + "."}

    app.include_router(router)
