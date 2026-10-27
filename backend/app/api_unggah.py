"""API unggah file data lewat web (Fase 6) + pratinjau sebelum masuk database.

Alur yang dipakai layar "Unggah Data" (v0.9.8):

1. ``POST /api/unggah/pratinjau`` — file diperiksa (divalidasi) **tanpa**
   menyimpan apa pun ke database. Semua baris yang bermasalah dikembalikan
   sebagai daftar *data mismatch*; transaksi di-rollback.
2. ``GET  /api/unggah/pratinjau/{token}/mismatch.xlsx`` — daftar baris
   bermasalah diunduh sebagai Excel (baris asli + alasan + lokasi), diperbaiki
   pengguna, lalu diunggah ulang.
3. ``POST /api/unggah/pratinjau/{token}/konfirmasi`` — baru di sini data
   benar-benar dimasukkan ke database.
4. ``DELETE /api/unggah/pratinjau/{token}`` — batalkan (file sementara dibuang).

``POST /api/unggah`` yang lama tetap ada (langsung impor tanpa pratinjau) untuk
skrip/otomatisasi yang sudah memakainya.

Aturan file: hanya .xlsx, batas ukuran UNGGAH_MAKS_MB (default 20 MB).
Format: ekspor "Master Data" (tiket intake) atau format template (sheet
Pembelian_Harian / Stok_Harian / SICOM_Harian / Kurs / Parameter_Pabrik /
Aturan_Stok_Harga / Target_Margin / Perencanaan / Master_Vendor / Master_Grade /
Intake_Tiket / Harga_Grade) — boleh sebagian sheet.
"""
from __future__ import annotations

import datetime as dt
import io
import os
import pathlib
import re
import secrets
import shutil
import threading
import time

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile
from sqlalchemy.orm import Session

from .core.alarm import catat_audit
from .importer.template import impor_master_data, impor_template
from .importer.tiket import cari_header

SHEET_TEMPLATE = {
    "Pembelian_Harian", "SICOM_Harian", "Kurs", "Stok_Harian",
    "Parameter_Pabrik", "Aturan_Stok_Harga", "Target_Margin", "Perencanaan",
    "Master_Vendor", "Master_Grade", "Intake_Tiket", "Harga_Grade",
}

BATAS_PERINGATAN = 200      # maksimal peringatan yang dikirim ke layar
BATAS_MISMATCH_LAYAR = 300  # daftar mismatch di layar; versi Excel lengkap
UMUR_PRATINJAU = 60 * 60    # detik: pratinjau kedaluwarsa setelah 1 jam
NAMA_JENIS = {
    "tiket": "Tiket intake (Master Data)",
    "template": "Data berkala (format template)",
}

#: Pratinjau yang menunggu konfirmasi (disimpan di memori proses).
_PRATINJAU: dict[str, dict] = {}
_KUNCI = threading.Lock()


def _nama_aman(nama: str) -> str:
    dasar = pathlib.Path(nama).name
    return re.sub(r"[^A-Za-z0-9._ -]", "_", dasar) or "unggahan.xlsx"


def _folder_unggah() -> pathlib.Path:
    folder = pathlib.Path(os.environ.get("UNGGAH_DIR", "unggahan"))
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _folder_pratinjau() -> pathlib.Path:
    folder = _folder_unggah() / ".pratinjau"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def deteksi_jenis(path: str) -> str | None:
    """'template' | 'tiket' | None (tak dikenal)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if set(wb.sheetnames) & SHEET_TEMPLATE:
            return "template"
        for nama in wb.sheetnames:
            if cari_header(wb[nama]) is not None:
                return "tiket"
        return None
    finally:
        wb.close()


# --------------------------------------------------------------- berkas masuk

async def _baca_file(file: UploadFile) -> tuple[str, bytes]:
    """Periksa nama, ekstensi, dan ukuran; kembalikan (nama_aman, isi)."""
    nama = _nama_aman(file.filename or "unggahan.xlsx")
    if not nama.lower().endswith(".xlsx"):
        raise HTTPException(422, f"File '{nama}' bukan .xlsx. Hanya file Excel "
                                 ".xlsx yang diterima — simpan/ekspor ulang "
                                 "sebagai .xlsx lalu unggah lagi.")
    isi = await file.read()
    maks_mb = float(os.environ.get("UNGGAH_MAKS_MB", "20"))
    if len(isi) > maks_mb * 1024 * 1024:
        raise HTTPException(413, f"File {len(isi) / 1024 / 1024:.1f} MB melebihi "
                                 f"batas {maks_mb:.0f} MB. Pecah filenya "
                                 "(mis. per tahun) lalu unggah bertahap.")
    if not isi:
        raise HTTPException(422, "File kosong. Pilih file yang benar.")
    return nama, isi


def _deteksi_atau_tolak(path: str, nama: str) -> str:
    """Tentukan jenis file; hapus berkas sementara bila ternyata tidak layak."""
    tujuan = pathlib.Path(path)

    def tolak(pesan: str) -> HTTPException:
        tujuan.unlink(missing_ok=True)
        return HTTPException(422, pesan)

    try:
        jenis = deteksi_jenis(path)
    except Exception:
        raise tolak(f"File '{nama}' tidak bisa dibuka sebagai Excel. Pastikan "
                    "file tidak rusak/terkunci password, lalu coba lagi.")
    if jenis is None:
        raise tolak(f"Format '{nama}' tidak dikenal. Yang diterima: (1) ekspor "
                    "'Master Data' tiket dari aplikasi internal, atau (2) file "
                    "dengan sheet format template: "
                    + ", ".join(sorted(SHEET_TEMPLATE)) + ".")
    return jenis


def _fungsi_impor(jenis: str):
    return impor_master_data if jenis == "tiket" else impor_template


# ------------------------------------------------------------- data mismatch

def _baris_ditolak(rincian: list[dict]) -> list[dict]:
    """Rapikan rincian jadi satu baris per (sheet, nomor baris).

    Satu baris bisa ditolak karena beberapa hal sekaligus; alasannya digabung
    supaya saat diperbaiki di Excel pengguna melihat semua penyebabnya.
    """
    hasil: list[dict] = []
    posisi: dict[tuple, int] = {}
    for r in rincian:
        if r.get("jenis") != "ditolak":
            continue
        kunci = (r.get("sheet"), r.get("baris"))
        alasan = str(r.get("pesan") or "").split(" Baris DITOLAK")[0].strip()
        kode = str(r.get("kode") or "")
        if kunci in posisi:
            lama = hasil[posisi[kunci]]
            if kode and kode not in lama["kode"].split("/"):
                lama["kode"] += "/" + kode
            if alasan and alasan not in lama["alasan"]:
                lama["alasan"] += " " + alasan
            continue
        posisi[kunci] = len(hasil)
        hasil.append({
            "sheet": r.get("sheet") or "-",
            "baris": r.get("baris"),
            "kode": kode or "PERIKSA",
            "alasan": alasan or "Baris tidak lolos pemeriksaan.",
        })
    return hasil


def _baris_header(ws) -> tuple[int, list]:
    """Tebak baris header: baris dengan sel terisi terbanyak di 10 baris awal."""
    terbaik, isi_terbaik = 1, -1
    for i, baris in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True),
                              start=1):
        isi = sum(1 for v in baris if v is not None and str(v).strip() != "")
        if isi > isi_terbaik:
            terbaik, isi_terbaik = i, isi
    return terbaik, [c.value for c in ws[terbaik]]


def _nama_sheet_aman(nama: str, dipakai: set[str]) -> str:
    dasar = re.sub(r"[\[\]:*?/\\]", "_", str(nama))[:28] or "Sheet"
    calon, n = dasar, 2
    while calon.lower() in dipakai:
        calon = f"{dasar[:26]}_{n}"
        n += 1
    dipakai.add(calon.lower())
    return calon


def buat_excel_mismatch(path_asli: str, nama_file: str,
                        mismatch: list[dict]) -> bytes:
    """Excel berisi baris asli yang ditolak + alasan + lokasi.

    Tiap sheet sumber dibuatkan satu lembar (header asli dipertahankan supaya
    file hasil perbaikan bisa diunggah ulang) dengan tiga kolom tambahan di
    kanan: KODE_MASALAH, ALASAN_DITOLAK, BARIS_DI_FILE_ASLI.
    """
    asal = openpyxl.load_workbook(path_asli, data_only=True)
    wb = openpyxl.Workbook()

    ringkas = wb.active
    ringkas.title = "Ringkasan"
    ringkas.append(["SHEET", "BARIS", "KODE MASALAH", "ALASAN (yang harus diperbaiki)"])
    for m in mismatch:
        ringkas.append([m["sheet"], m["baris"] or "-", m["kode"], m["alasan"]])
    if not mismatch:
        ringkas.append(["-", "-", "-", "Tidak ada baris bermasalah."])

    tebal = openpyxl.styles.Font(bold=True)
    merah = openpyxl.styles.PatternFill("solid", fgColor="FFE1E1")
    bingkai = openpyxl.styles.Alignment(vertical="top", wrap_text=True)

    for sel in ringkas[1]:
        sel.font, sel.fill = tebal, merah
    for kolom, lebar in zip("ABCD", (24, 10, 22, 70)):
        ringkas.column_dimensions[kolom].width = lebar
    ringkas.freeze_panes = "A2"

    kelompok: dict[str, list[dict]] = {}
    for m in mismatch:
        kelompok.setdefault(m["sheet"], []).append(m)

    dipakai: set[str] = {"ringkasan"}
    for nama_sheet, daftar in kelompok.items():
        if nama_sheet not in asal.sheetnames:
            continue  # lokasi tanpa baris (mis. aturan umum) — cukup di Ringkasan
        ws_asal = asal[nama_sheet]
        baris_header, judul = _baris_header(ws_asal)
        lebar_asli = max(len(judul), ws_asal.max_column or 0)
        judul = list(judul) + [None] * (lebar_asli - len(judul))
        ws = wb.create_sheet(_nama_sheet_aman(nama_sheet, dipakai))
        ws.append([("" if v is None else v) for v in judul]
                  + ["KODE_MASALAH", "ALASAN_DITOLAK", "BARIS_DI_FILE_ASLI"])
        for sel in ws[1]:
            sel.font, sel.fill = tebal, merah
        for m in daftar:
            nomor = m["baris"]
            if not nomor:
                continue
            nilai = [ws_asal.cell(row=nomor, column=k + 1).value
                     for k in range(lebar_asli)]
            ws.append([("" if v is None else v) for v in nilai]
                      + [m["kode"], m["alasan"], nomor])
        for sel in ws[1]:
            ws.column_dimensions[sel.column_letter].width = max(
                12, min(40, len(str(sel.value or "")) + 4))
        for baris in ws.iter_rows(min_row=2):
            baris[-3].alignment = baris[-2].alignment = bingkai
        ws.freeze_panes = "A2"

    # Lembar petunjuk singkat
    bantu = wb.create_sheet("Cara memperbaiki")
    for teks in [
        f"File asal: {nama_file}",
        "Baris di bawah ini BELUM masuk database.",
        "1) Buka sheet yang sesuai, periksa kolom ALASAN_DITOLAK.",
        "2) Perbaiki nilai pada baris tersebut (jangan mengubah judul kolom).",
        "3) Hapus kolom KODE_MASALAH, ALASAN_DITOLAK, dan BARIS_DI_FILE_ASLI.",
        "4) Unggah ulang file ini lewat menu Unggah Data.",
        "Baris yang tidak bermasalah tidak perlu diunggah ulang.",
    ]:
        bantu.append([teks])
    bantu.column_dimensions["A"].width = 80

    aliran = io.BytesIO()
    wb.save(aliran)
    asal.close()
    return aliran.getvalue()


# -------------------------------------------------------------- simpan-token

def _bersihkan_pratinjau() -> None:
    """Buang pratinjau yang kedaluwarsa (beserta file sementaranya)."""
    batas = time.time() - UMUR_PRATINJAU
    with _KUNCI:
        kedaluwarsa = [t for t, d in _PRATINJAU.items() if d["waktu"] < batas]
        data_kedaluwarsa = [_PRATINJAU.pop(t) for t in kedaluwarsa]
    for data in data_kedaluwarsa:
        if data.get("sementara"):
            try:
                pathlib.Path(data["path"]).unlink(missing_ok=True)
            except OSError:
                pass


def _ambil_pratinjau(token: str) -> dict:
    _bersihkan_pratinjau()
    with _KUNCI:
        data = _PRATINJAU.get(token)
    if data is None:
        raise HTTPException(410, "Pratinjau sudah kedaluwarsa atau sudah "
                                 "dipakai. Silakan unggah ulang filenya.")
    if not pathlib.Path(data["path"]).exists():
        with _KUNCI:
            _PRATINJAU.pop(token, None)
        raise HTTPException(410, "File pratinjau sudah tidak ada. Silakan "
                                 "unggah ulang filenya.")
    return data


def _buang_pratinjau(token: str) -> None:
    """Lupakan pratinjau; file sementara dibuang, file yang sudah dikonfirmasi
    (kini tersimpan permanen) tidak disentuh."""
    with _KUNCI:
        data = _PRATINJAU.pop(token, None)
    if data and data.get("sementara"):
        try:
            pathlib.Path(data["path"]).unlink(missing_ok=True)
        except OSError:
            pass


def _stempel() -> str:
    return dt.datetime.now().strftime("%Y%m%d_%H%M%S")


def pasang_router_unggah(app, ambil_session):
    router = APIRouter(prefix="/api")
    dep = Depends(ambil_session)

    # ------------------------------------------------------------- pratinjau
    @router.post("/unggah/pratinjau")
    async def pratinjau(file: UploadFile = File(...), pengisi: str = Form("?"),
                        session: Session = dep):
        nama, isi = await _baca_file(file)
        _bersihkan_pratinjau()

        token = secrets.token_urlsafe(18)
        path = _folder_pratinjau() / f"{token}.xlsx"
        path.write_bytes(isi)
        jenis = _deteksi_atau_tolak(str(path), nama)

        langkah: list[str] = []
        fn = _fungsi_impor(jenis)
        hasil = None
        try:
            hasil = fn(str(path), session, pengunggah=pengisi,
                       simpan=False, kemajuan=langkah.append)
        except HTTPException:
            raise
        except Exception as e:  # file rusak / isi tak terduga
            session.rollback()
            path.unlink(missing_ok=True)
            raise HTTPException(422, f"File '{nama}' gagal diperiksa: {e}. "
                                     "Perbaiki file lalu unggah ulang.")
        finally:
            # Tidak satu baris pun boleh tersimpan pada tahap pratinjau.
            session.rollback()

        mismatch = _baris_ditolak(hasil.rincian)
        jumlah = dict(hasil.jumlah)
        ditolak = dict(hasil.ditolak)
        _PRATINJAU[token] = {
            "waktu": time.time(), "path": str(path), "nama_file": nama,
            "jenis": jenis, "pengisi": pengisi, "jumlah": jumlah,
            "ditolak": ditolak, "mismatch": mismatch,
            "sementara": True, "dikonfirmasi": False,
        }

        total_ok = sum(v for v in jumlah.values() if isinstance(v, int))
        total_tolak = sum(v for v in ditolak.values() if isinstance(v, int))
        if total_tolak:
            pesan = (f"Pratinjau siap: {total_ok} baris siap disimpan, "
                     f"{total_tolak} baris bermasalah dan TIDAK akan dimasukkan. "
                     "Unduh daftar mismatch bila perlu diperbaiki, lalu tekan "
                     "Konfirmasi simpan.")
        else:
            pesan = (f"Pratinjau siap: {total_ok} baris siap disimpan, tidak ada "
                     "baris bermasalah. Tekan Konfirmasi simpan untuk "
                     "memasukkan ke database.")

        return {
            "status": "pratinjau",
            "token": token,
            "jenis": NAMA_JENIS[jenis],
            "nama_file": nama,
            "duplikat_file": hasil.duplikat_file,
            "jumlah": jumlah,
            "total_jumlah": total_ok,
            "ditolak": ditolak,
            "total_ditolak": total_tolak,
            "mismatch": mismatch[:BATAS_MISMATCH_LAYAR],
            "jumlah_mismatch": len(mismatch),
            "jumlah_peringatan": len(hasil.peringatan),
            "peringatan": hasil.peringatan[:BATAS_PERINGATAN],
            "langkah": langkah,
            "kedaluwarsa_menit": UMUR_PRATINJAU // 60,
            "pesan": pesan,
        }

    # ----------------------------------------------- unduh Excel mismatch
    @router.get("/unggah/pratinjau/{token}/mismatch.xlsx")
    def unduh_mismatch(token: str):
        data = _ambil_pratinjau(token)
        mismatch = data.get("mismatch") or []
        if not mismatch:
            raise HTTPException(404, "Tidak ada baris mismatch pada pratinjau ini.")
        isi = buat_excel_mismatch(data["path"], data["nama_file"], mismatch)
        nama = f"mismatch_{_stempel()}.xlsx"
        return Response(
            content=isi,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{nama}"'},
        )

    # ------------------------------------------------------------ konfirmasi
    @router.post("/unggah/pratinjau/{token}/konfirmasi")
    def konfirmasi(token: str, session: Session = dep):
        data = _ambil_pratinjau(token)
        if data.get("dikonfirmasi"):
            raise HTTPException(409, "Pratinjau ini sudah dikonfirmasi dan datanya "
                                     "sudah masuk database. Tidak disimpan dua kali.")
        nama, jenis = data["nama_file"], data["jenis"]

        # Baru sekarang file dipindah ke folder unggahan tetap (ikut backup).
        tujuan = _folder_unggah() / f"{_stempel()}_{nama}"
        asal = pathlib.Path(data["path"])
        try:
            shutil.move(str(asal), str(tujuan))
        except PermissionError as e:
            raise HTTPException(
                423,
                "File pratinjau masih sedang dipakai Windows. Tutup file itu "
                "bila terbuka di Excel atau aplikasi lain, tunggu sebentar, "
                "lalu tekan Konfirmasi simpan lagi.",
            ) from e

        fn = _fungsi_impor(jenis)
        try:
            hasil = fn(str(tujuan), session, pengunggah=data["pengisi"])
        except Exception as e:
            session.rollback()
            try:
                shutil.move(str(tujuan), str(asal))
            except OSError:
                pass
            raise HTTPException(422, f"Data dari '{nama}' gagal disimpan: {e}. "
                                     "Tidak ada perubahan pada database; "
                                     "silakan periksa kembali lalu coba lagi.")

        catat_audit(session, data["pengisi"], "unggah", "file", {
            "nama_file": nama, "jenis": jenis, "unggahan_id": hasil.unggahan_id,
            "jumlah": hasil.jumlah, "ditolak": hasil.ditolak,
            "duplikat_file": hasil.duplikat_file, "lewat_pratinjau": True,
        })
        session.commit()

        # File sudah menjadi arsip permanen; pratinjau tetap sebentar disimpan
        # supaya daftar mismatch masih bisa diunduh untuk diperbaiki.
        mismatch = _baris_ditolak(hasil.rincian)
        with _KUNCI:
            _PRATINJAU[token] = {**data, "path": str(tujuan), "sementara": False,
                                 "dikonfirmasi": True, "waktu": time.time(),
                                 "mismatch": mismatch}

        return {
            "status": "selesai",
            "token": token,
            "mismatch_tersedia": bool(mismatch),
            "jenis": NAMA_JENIS[jenis],
            "nama_file": nama,
            "unggahan_id": hasil.unggahan_id,
            "duplikat_file": hasil.duplikat_file,
            "jumlah": hasil.jumlah,
            "ditolak": hasil.ditolak,
            "total_ditolak": sum(v for v in hasil.ditolak.values()
                                 if isinstance(v, int)),
            "jumlah_mismatch": len(mismatch),
            "mismatch": mismatch[:BATAS_MISMATCH_LAYAR],
            "jumlah_peringatan": len(hasil.peringatan),
            "peringatan": hasil.peringatan[:BATAS_PERINGATAN],
            "pesan": ("File sama persis pernah diunggah — data tidak digandakan "
                      "(aman)." if hasil.duplikat_file else
                      "Data sudah masuk database; mart & alarm sudah dihitung "
                      "ulang otomatis."),
        }

    # -------------------------------------------------------------- batalkan
    @router.delete("/unggah/pratinjau/{token}")
    def batalkan(token: str):
        _buang_pratinjau(token)
        return {"status": "dibatalkan",
                "pesan": "Pratinjau dibatalkan; tidak ada data yang disimpan."}

    # ------------------------------------------- unggah langsung (kompatibel)
    @router.post("/unggah")
    async def unggah(file: UploadFile = File(...), pengisi: str = Form("?"),
                     session: Session = dep):
        nama, isi = await _baca_file(file)

        # Simpan file asli apa adanya (ikut backup harian; tidak pernah dihapus)
        path = _folder_unggah() / f"{_stempel()}_{nama}"
        path.write_bytes(isi)
        jenis = _deteksi_atau_tolak(str(path), nama)

        fn = _fungsi_impor(jenis)
        hasil = fn(str(path), session, pengunggah=pengisi)

        catat_audit(session, pengisi, "unggah", "file", {
            "nama_file": nama, "jenis": jenis, "ukuran_kb": round(len(isi) / 1024, 1),
            "unggahan_id": hasil.unggahan_id, "jumlah": hasil.jumlah,
            "ditolak": hasil.ditolak, "duplikat_file": hasil.duplikat_file,
        })
        session.commit()

        return {
            "status": "selesai",
            "jenis": NAMA_JENIS[jenis],
            "nama_file": nama,
            "unggahan_id": hasil.unggahan_id,
            "duplikat_file": hasil.duplikat_file,
            "jumlah": hasil.jumlah,
            "ditolak": hasil.ditolak,
            "jumlah_peringatan": len(hasil.peringatan),
            "peringatan": hasil.peringatan[:BATAS_PERINGATAN],
            "pesan": ("File sama persis pernah diunggah — data tidak digandakan "
                      "(aman)." if hasil.duplikat_file else
                      "Impor selesai; mart & alarm sudah dihitung ulang otomatis."),
        }

    app.include_router(router)
